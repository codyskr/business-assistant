import asyncio
import json
import logging
import mimetypes
import os
import sqlite3
import tempfile
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import httpx
from apscheduler.schedulers.asyncio import AsyncIOScheduler
from cryptography.fernet import Fernet
from dotenv import load_dotenv
from assistant_core.commands import (
    build_knowledge_reply,
    build_notes_reply,
    build_reminders_reply,
    resolve_help_section,
)
from telegram import InlineKeyboardButton, InlineKeyboardMarkup, Update
from telegram.constants import ChatAction
from telegram.ext import (
    Application,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
    MessageHandler,
    filters,
)


load_dotenv()

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
DB_PATH = DATA_DIR / "agent.db"
REMINDERS_MD_PATH = DATA_DIR / "reminders.md"
NOTES_PATH = DATA_DIR / "notes.jsonl"
KNOWLEDGE_PATH = DATA_DIR / "knowledge.jsonl"
CREDENTIALS_PATH = BASE_DIR / "credentials.json"
TOKEN_PATH = BASE_DIR / "token.json"
GOOGLE_TOKENS_DIR = DATA_DIR / "google_tokens"

TELEGRAM_BOT_TOKEN = os.getenv("TELEGRAM_BOT_TOKEN")
ALLOWED_TELEGRAM_USER_IDS = {
    int(value.strip())
    for value in os.getenv("ALLOWED_TELEGRAM_USER_IDS", "").split(",")
    if value.strip().isdigit()
}

OLLAMA_URL = os.getenv("OLLAMA_URL", "http://127.0.0.1:11434")
OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "llama3.2:3b")
NOTES_LLM_SEARCH_ENABLED = os.getenv("NOTES_LLM_SEARCH_ENABLED", "true").lower() in {
    "1",
    "true",
    "yes",
    "on",
}
NOTES_LLM_SEARCH_LIMIT = int(os.getenv("NOTES_LLM_SEARCH_LIMIT", "30"))
KNOWLEDGE_LLM_SEARCH_ENABLED = os.getenv("KNOWLEDGE_LLM_SEARCH_ENABLED", "true").lower() in {
    "1",
    "true",
    "yes",
    "on",
}
KNOWLEDGE_LLM_SEARCH_LIMIT = int(os.getenv("KNOWLEDGE_LLM_SEARCH_LIMIT", "40"))
KNOWLEDGE_MAX_FILE_MB = int(os.getenv("KNOWLEDGE_MAX_FILE_MB", "20"))
KNOWLEDGE_CHUNK_CHARS = int(os.getenv("KNOWLEDGE_CHUNK_CHARS", "1400"))
KNOWLEDGE_CHUNK_OVERLAP = int(os.getenv("KNOWLEDGE_CHUNK_OVERLAP", "180"))
TIMEZONE = os.getenv("TIMEZONE", "Europe/Moscow")
PRIVACY_MODE = os.getenv("PRIVACY_MODE", "balanced").strip().lower()
STT_PROVIDER = os.getenv("STT_PROVIDER", "local").strip().lower()
STT_FALLBACK_PROVIDER = os.getenv("STT_FALLBACK_PROVIDER", "local").strip().lower()
STT_MAX_VOICE_SECONDS = int(os.getenv("STT_MAX_VOICE_SECONDS", "300"))
WHISPER_MODEL = os.getenv("WHISPER_MODEL", "tiny")
GROQ_API_KEY = os.getenv("GROQ_API_KEY", "")
GROQ_STT_MODEL = os.getenv("GROQ_STT_MODEL", "whisper-large-v3-turbo")
GROQ_STT_URL = os.getenv(
    "GROQ_STT_URL",
    "https://api.groq.com/openai/v1/audio/transcriptions",
)
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY", "")
OPENAI_STT_MODEL = os.getenv("OPENAI_STT_MODEL", "gpt-4o-mini-transcribe")
OPENAI_STT_URL = os.getenv(
    "OPENAI_STT_URL",
    "https://api.openai.com/v1/audio/transcriptions",
)
GOOGLE_CALENDAR_ID = os.getenv("GOOGLE_CALENDAR_ID", "primary")
AUTO_CONFIRM_ACTIONS = os.getenv("AUTO_CONFIRM_ACTIONS", "false").lower() in {
    "1",
    "true",
    "yes",
    "on",
}
DEFAULT_EVENT_REMINDER_MINUTES = int(os.getenv("DEFAULT_EVENT_REMINDER_MINUTES", "15"))
MEMORY_ENABLED = os.getenv("MEMORY_ENABLED", "true").lower() in {
    "1",
    "true",
    "yes",
    "on",
}
MEMORY_KEY_PATH = Path(os.getenv("MEMORY_KEY_PATH", str(DATA_DIR / "memory.key")))
MEMORY_RETENTION_DAYS = int(os.getenv("MEMORY_RETENTION_DAYS", "7"))
MEMORY_RECENT_MESSAGES = int(os.getenv("MEMORY_RECENT_MESSAGES", "12"))
MEMORY_RECENT_ACTIONS = int(os.getenv("MEMORY_RECENT_ACTIONS", "6"))

GOOGLE_SCOPES = ["https://www.googleapis.com/auth/calendar.events"]

logging.basicConfig(
    format="%(asctime)s %(levelname)s %(name)s: %(message)s",
    level=logging.INFO,
)
logger = logging.getLogger(__name__)

try:
    from rapidfuzz import fuzz
except ImportError:
    fuzz = None


@dataclass
class ParsedIntent:
    kind: str
    data: dict[str, Any]
    reply: str


class Store:
    def __init__(self, path: Path) -> None:
        self.path = path
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._init()

    def connect(self) -> sqlite3.Connection:
        conn = sqlite3.connect(self.path)
        conn.row_factory = sqlite3.Row
        return conn

    def _init(self) -> None:
        with self.connect() as conn:
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS pending_actions (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    chat_id INTEGER NOT NULL,
                    user_id INTEGER NOT NULL,
                    action_json TEXT NOT NULL,
                    created_at TEXT NOT NULL
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS reminders (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    chat_id INTEGER NOT NULL,
                    text TEXT NOT NULL,
                    remind_at TEXT NOT NULL,
                    status TEXT NOT NULL DEFAULT 'pending',
                    created_at TEXT NOT NULL
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS memory_messages (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    chat_id INTEGER NOT NULL,
                    user_id INTEGER,
                    role TEXT NOT NULL,
                    text_cipher TEXT NOT NULL,
                    created_at TEXT NOT NULL
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS action_history (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    chat_id INTEGER NOT NULL,
                    kind TEXT NOT NULL,
                    summary_cipher TEXT NOT NULL,
                    payload_cipher TEXT NOT NULL,
                    created_at TEXT NOT NULL
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS event_selections (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    chat_id INTEGER NOT NULL,
                    action TEXT NOT NULL,
                    event_json TEXT NOT NULL,
                    payload_json TEXT NOT NULL,
                    created_at TEXT NOT NULL
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS reminder_selections (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    chat_id INTEGER NOT NULL,
                    reminder_ids_json TEXT NOT NULL,
                    created_at TEXT NOT NULL
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS debug_log (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    chat_id INTEGER NOT NULL,
                    user_id INTEGER,
                    source TEXT NOT NULL,
                    stt_provider TEXT,
                    input_cipher TEXT NOT NULL,
                    intent_kind TEXT,
                    intent_cipher TEXT,
                    result TEXT,
                    error TEXT,
                    created_at TEXT NOT NULL
                )
                """
            )
            conn.execute(
                """
                CREATE TABLE IF NOT EXISTS user_settings (
                    user_id INTEGER PRIMARY KEY,
                    google_calendar_id TEXT,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                )
                """
            )

    def add_pending(self, chat_id: int, user_id: int, action: dict[str, Any]) -> int:
        with self.connect() as conn:
            cur = conn.execute(
                """
                INSERT INTO pending_actions (chat_id, user_id, action_json, created_at)
                VALUES (?, ?, ?, ?)
                """,
                (chat_id, user_id, json.dumps(action, ensure_ascii=False), now_iso()),
            )
            return int(cur.lastrowid)

    def pop_pending(self, action_id: int) -> sqlite3.Row | None:
        with self.connect() as conn:
            row = conn.execute(
                "SELECT * FROM pending_actions WHERE id = ?",
                (action_id,),
            ).fetchone()
            if row:
                conn.execute("DELETE FROM pending_actions WHERE id = ?", (action_id,))
            return row

    def add_reminder(self, chat_id: int, text: str, remind_at: str) -> int:
        with self.connect() as conn:
            cur = conn.execute(
                """
                INSERT INTO reminders (chat_id, text, remind_at, status, created_at)
                VALUES (?, ?, ?, 'pending', ?)
                """,
                (chat_id, text, remind_at, now_iso()),
            )
            return int(cur.lastrowid)

    def list_pending_reminders(self) -> list[sqlite3.Row]:
        with self.connect() as conn:
            return list(
                conn.execute(
                    "SELECT * FROM reminders WHERE status = 'pending' ORDER BY remind_at"
                ).fetchall()
            )

    def list_reminders_between(self, chat_id: int, start: datetime, end: datetime) -> list[sqlite3.Row]:
        with self.connect() as conn:
            return list(
                conn.execute(
                    """
                    SELECT *
                    FROM reminders
                    WHERE chat_id = ?
                      AND status = 'pending'
                      AND remind_at >= ?
                      AND remind_at < ?
                    ORDER BY remind_at
                    """,
                    (
                        chat_id,
                        start.isoformat(timespec="seconds"),
                        end.isoformat(timespec="seconds"),
                    ),
                ).fetchall()
            )

    def list_pending_reminders_for_chat(self, chat_id: int, limit: int = 500) -> list[sqlite3.Row]:
        with self.connect() as conn:
            return list(
                conn.execute(
                    """
                    SELECT *
                    FROM reminders
                    WHERE chat_id = ?
                      AND status = 'pending'
                    ORDER BY remind_at
                    LIMIT ?
                    """,
                    (chat_id, limit),
                ).fetchall()
            )

    def mark_reminder_sent(self, reminder_id: int) -> None:
        with self.connect() as conn:
            conn.execute(
                "UPDATE reminders SET status = 'sent' WHERE id = ?",
                (reminder_id,),
            )

    def cancel_reminders(self, chat_id: int, reminder_ids: list[int]) -> int:
        if not reminder_ids:
            return 0
        placeholders = ",".join("?" for _ in reminder_ids)
        with self.connect() as conn:
            cur = conn.execute(
                f"""
                UPDATE reminders
                SET status = 'cancelled'
                WHERE chat_id = ?
                  AND status = 'pending'
                  AND id IN ({placeholders})
                """,
                [chat_id, *reminder_ids],
            )
            return int(cur.rowcount)

    def search_reminders(
        self,
        chat_id: int,
        query: str | None,
        start: datetime | None = None,
        end: datetime | None = None,
        limit: int = 100,
    ) -> list[sqlite3.Row]:
        clauses = ["chat_id = ?", "status = 'pending'"]
        values: list[Any] = [chat_id]
        if start is not None:
            clauses.append("remind_at >= ?")
            values.append(start.isoformat(timespec="seconds"))
        if end is not None:
            clauses.append("remind_at < ?")
            values.append(end.isoformat(timespec="seconds"))
        if query:
            clauses.append("LOWER(text) LIKE ?")
            values.append(f"%{query.lower()}%")
        values.append(limit)
        with self.connect() as conn:
            return list(
                conn.execute(
                    f"""
                    SELECT *
                    FROM reminders
                    WHERE {' AND '.join(clauses)}
                    ORDER BY remind_at
                    LIMIT ?
                    """,
                    values,
                ).fetchall()
            )

    def add_memory_message(self, chat_id: int, user_id: int | None, role: str, text: str) -> None:
        if not MEMORY_ENABLED or not text.strip():
            return
        with self.connect() as conn:
            conn.execute(
                """
                INSERT INTO memory_messages (chat_id, user_id, role, text_cipher, created_at)
                VALUES (?, ?, ?, ?, ?)
                """,
                (chat_id, user_id, role, encrypt_text(text), now_iso()),
            )

    def recent_memory_messages(self, chat_id: int, limit: int) -> list[sqlite3.Row]:
        cutoff = datetime.now(ZoneInfo(TIMEZONE)) - timedelta(days=MEMORY_RETENTION_DAYS)
        with self.connect() as conn:
            return list(
                conn.execute(
                    """
                    SELECT role, text_cipher, created_at
                    FROM memory_messages
                    WHERE chat_id = ? AND created_at >= ?
                    ORDER BY id DESC
                    LIMIT ?
                    """,
                    (chat_id, cutoff.isoformat(timespec="seconds"), limit),
                ).fetchall()
            )

    def add_action_history(self, chat_id: int, kind: str, summary: str, payload: dict[str, Any]) -> None:
        if not MEMORY_ENABLED:
            return
        with self.connect() as conn:
            conn.execute(
                """
                INSERT INTO action_history (chat_id, kind, summary_cipher, payload_cipher, created_at)
                VALUES (?, ?, ?, ?, ?)
                """,
                (
                    chat_id,
                    kind,
                    encrypt_text(summary),
                    encrypt_text(json.dumps(payload, ensure_ascii=False)),
                    now_iso(),
                ),
            )

    def recent_action_history(self, chat_id: int, limit: int) -> list[sqlite3.Row]:
        cutoff = datetime.now(ZoneInfo(TIMEZONE)) - timedelta(days=MEMORY_RETENTION_DAYS)
        with self.connect() as conn:
            return list(
                conn.execute(
                    """
                    SELECT kind, summary_cipher, created_at
                    FROM action_history
                    WHERE chat_id = ? AND created_at >= ?
                    ORDER BY id DESC
                    LIMIT ?
                    """,
                    (chat_id, cutoff.isoformat(timespec="seconds"), limit),
                ).fetchall()
            )

    def recent_actions_for_display(self, chat_id: int, limit: int) -> list[sqlite3.Row]:
        with self.connect() as conn:
            return list(
                conn.execute(
                    """
                    SELECT kind, summary_cipher, created_at
                    FROM action_history
                    WHERE chat_id = ?
                    ORDER BY id DESC
                    LIMIT ?
                    """,
                    (chat_id, limit),
                ).fetchall()
            )

    def add_debug_log(
        self,
        *,
        chat_id: int,
        user_id: int | None,
        source: str,
        stt_provider: str | None,
        input_text: str,
        intent_kind: str | None = None,
        intent_payload: dict[str, Any] | None = None,
        result: str | None = None,
        error: str | None = None,
    ) -> int:
        with self.connect() as conn:
            cur = conn.execute(
                """
                INSERT INTO debug_log (
                    chat_id,
                    user_id,
                    source,
                    stt_provider,
                    input_cipher,
                    intent_kind,
                    intent_cipher,
                    result,
                    error,
                    created_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    chat_id,
                    user_id,
                    source,
                    stt_provider,
                    encrypt_text(input_text),
                    intent_kind,
                    encrypt_text(json.dumps(intent_payload, ensure_ascii=False))
                    if intent_payload is not None
                    else None,
                    result,
                    error,
                    now_iso(),
                ),
            )
            return int(cur.lastrowid)

    def update_debug_log(
        self,
        debug_id: int | None,
        *,
        intent_kind: str | None = None,
        intent_payload: dict[str, Any] | None = None,
        result: str | None = None,
        error: str | None = None,
    ) -> None:
        if debug_id is None:
            return

        updates: list[str] = []
        values: list[Any] = []
        if intent_kind is not None:
            updates.append("intent_kind = ?")
            values.append(intent_kind)
        if intent_payload is not None:
            updates.append("intent_cipher = ?")
            values.append(encrypt_text(json.dumps(intent_payload, ensure_ascii=False)))
        if result is not None:
            updates.append("result = ?")
            values.append(result)
        if error is not None:
            updates.append("error = ?")
            values.append(error)
        if not updates:
            return

        values.append(debug_id)
        with self.connect() as conn:
            conn.execute(
                f"UPDATE debug_log SET {', '.join(updates)} WHERE id = ?",
                values,
            )

    def recent_debug_logs(self, chat_id: int, limit: int) -> list[sqlite3.Row]:
        with self.connect() as conn:
            return list(
                conn.execute(
                    """
                    SELECT *
                    FROM debug_log
                    WHERE chat_id = ?
                    ORDER BY id DESC
                    LIMIT ?
                    """,
                    (chat_id, limit),
                ).fetchall()
            )

    def count_memory_messages(self, chat_id: int) -> int:
        with self.connect() as conn:
            row = conn.execute(
                "SELECT COUNT(*) AS total FROM memory_messages WHERE chat_id = ?",
                (chat_id,),
            ).fetchone()
            return int(row["total"])

    def forget_today(self, chat_id: int) -> int:
        start = datetime.now(ZoneInfo(TIMEZONE)).replace(
            hour=0,
            minute=0,
            second=0,
            microsecond=0,
        )
        with self.connect() as conn:
            cur = conn.execute(
                "DELETE FROM memory_messages WHERE chat_id = ? AND created_at >= ?",
                (chat_id, start.isoformat(timespec="seconds")),
            )
            return int(cur.rowcount)

    def add_event_selection(
        self,
        chat_id: int,
        action: str,
        event: dict[str, Any],
        payload: dict[str, Any],
    ) -> int:
        with self.connect() as conn:
            cur = conn.execute(
                """
                INSERT INTO event_selections (chat_id, action, event_json, payload_json, created_at)
                VALUES (?, ?, ?, ?, ?)
                """,
                (
                    chat_id,
                    action,
                    json.dumps(event, ensure_ascii=False),
                    json.dumps(payload, ensure_ascii=False),
                    now_iso(),
                ),
            )
            return int(cur.lastrowid)

    def pop_event_selection(self, selection_id: int) -> sqlite3.Row | None:
        with self.connect() as conn:
            row = conn.execute(
                "SELECT * FROM event_selections WHERE id = ?",
                (selection_id,),
            ).fetchone()
            if row:
                conn.execute("DELETE FROM event_selections WHERE id = ?", (selection_id,))
            return row

    def add_reminder_selection(self, chat_id: int, reminder_ids: list[int]) -> int:
        with self.connect() as conn:
            cur = conn.execute(
                """
                INSERT INTO reminder_selections (chat_id, reminder_ids_json, created_at)
                VALUES (?, ?, ?)
                """,
                (chat_id, json.dumps(reminder_ids), now_iso()),
            )
            return int(cur.lastrowid)

    def pop_reminder_selection(self, selection_id: int) -> sqlite3.Row | None:
        with self.connect() as conn:
            row = conn.execute(
                "SELECT * FROM reminder_selections WHERE id = ?",
                (selection_id,),
            ).fetchone()
            if row:
                conn.execute("DELETE FROM reminder_selections WHERE id = ?", (selection_id,))
            return row

    def get_user_calendar_id(self, user_id: int) -> str:
        with self.connect() as conn:
            row = conn.execute(
                "SELECT google_calendar_id FROM user_settings WHERE user_id = ?",
                (user_id,),
            ).fetchone()
            if row and row["google_calendar_id"]:
                return str(row["google_calendar_id"])
        return GOOGLE_CALENDAR_ID

    def set_user_calendar_id(self, user_id: int, calendar_id: str) -> None:
        now = now_iso()
        with self.connect() as conn:
            conn.execute(
                """
                INSERT INTO user_settings (user_id, google_calendar_id, created_at, updated_at)
                VALUES (?, ?, ?, ?)
                ON CONFLICT(user_id) DO UPDATE SET
                    google_calendar_id = excluded.google_calendar_id,
                    updated_at = excluded.updated_at
                """,
                (user_id, calendar_id, now, now),
            )


def get_fernet() -> Fernet:
    MEMORY_KEY_PATH.parent.mkdir(parents=True, exist_ok=True)
    if MEMORY_KEY_PATH.exists():
        key = MEMORY_KEY_PATH.read_bytes().strip()
    else:
        key = Fernet.generate_key()
        MEMORY_KEY_PATH.write_bytes(key)
    return Fernet(key)


_fernet: Fernet | None = None


def encrypt_text(text: str) -> str:
    global _fernet
    if _fernet is None:
        _fernet = get_fernet()
    return _fernet.encrypt(text.encode("utf-8")).decode("ascii")


def decrypt_text(cipher: str) -> str:
    global _fernet
    if _fernet is None:
        _fernet = get_fernet()
    return _fernet.decrypt(cipher.encode("ascii")).decode("utf-8")


def normalize_search_text(text: str) -> str:
    import re

    return " ".join(re.findall(r"[a-zA-Zа-яА-ЯёЁ0-9]+", text.lower()))


def note_score(query: str, note_text: str) -> int:
    query_norm = normalize_search_text(query)
    text_norm = normalize_search_text(note_text)
    if not query_norm:
        return 0
    if not text_norm:
        return 0
    if fuzz is not None:
        return int(max(fuzz.partial_ratio(query_norm, text_norm), fuzz.token_set_ratio(query_norm, text_norm)))

    query_words = set(query_norm.split())
    text_words = set(text_norm.split())
    if not query_words:
        return 0
    return int(len(query_words & text_words) / len(query_words) * 100)


def expand_note_query(query: str) -> str:
    normalized = normalize_search_text(query)
    aliases = {
        "звон": "звонок созвон позвонить",
        "созвон": "звонок звонить встреча",
        "встреч": "созвон звонок мероприятие",
        "клиент": "заказчик контакт человек",
        "документ": "файл договор отчет pdf скан",
        "отчет": "документ файл сводка",
        "деньг": "оплата платеж счет сумма",
        "оплат": "деньги платеж счет сумма",
        "машин": "авто автомобиль",
        "авто": "машина автомобиль",
    }
    extra = []
    for word in normalized.split():
        for key, value in aliases.items():
            if word.startswith(key) or key.startswith(word):
                extra.append(value)
    return " ".join([query, *extra]).strip()


def add_note(chat_id: int, user_id: int | None, text: str) -> dict[str, Any]:
    NOTES_PATH.parent.mkdir(parents=True, exist_ok=True)
    note = {
        "id": f"note-{int(datetime.now(ZoneInfo(TIMEZONE)).timestamp() * 1000)}",
        "chat_id": chat_id,
        "user_id": user_id,
        "text": text.strip(),
        "created_at": now_iso(),
    }
    line = json.dumps(
        {
            "v": 1,
            "cipher": encrypt_text(json.dumps(note, ensure_ascii=False)),
        },
        ensure_ascii=False,
    )
    with NOTES_PATH.open("a", encoding="utf-8") as file:
        file.write(line + "\n")
    return note


def iter_notes(chat_id: int, user_id: int | None = None) -> list[dict[str, Any]]:
    if not NOTES_PATH.exists():
        return []

    notes: list[dict[str, Any]] = []
    with NOTES_PATH.open("r", encoding="utf-8") as file:
        for line in file:
            line = line.strip()
            if not line:
                continue
            try:
                wrapper = json.loads(line)
                note = json.loads(decrypt_text(str(wrapper["cipher"])))
            except Exception:
                logger.warning("Could not decrypt note line", exc_info=True)
                continue
            if int(note.get("chat_id", 0)) != chat_id:
                continue
            if user_id is not None and note.get("user_id") not in (None, user_id):
                continue
            notes.append(note)
    return notes


def search_notes(chat_id: int, query: str, user_id: int | None = None, limit: int = 5) -> list[dict[str, Any]]:
    expanded_query = expand_note_query(query)
    ranked = []
    for note in iter_notes(chat_id, user_id):
        score = note_score(expanded_query, str(note.get("text", "")))
        if score >= 25:
            ranked.append((score, note))
    ranked.sort(key=lambda item: (item[0], item[1].get("created_at", "")), reverse=True)
    return [note | {"score": score} for score, note in ranked[:limit]]


async def rerank_notes_with_ollama(query: str, notes: list[dict[str, Any]], limit: int) -> list[dict[str, Any]]:
    if not NOTES_LLM_SEARCH_ENABLED or not notes:
        return notes[:limit]

    numbered = []
    for index, note in enumerate(notes, start=1):
        numbered.append(
            {
                "index": index,
                "created_at": note.get("created_at"),
                "text": str(note.get("text", ""))[:900],
            }
        )

    system_prompt = """
You rerank encrypted local notes after they have been decrypted locally.
Return JSON only. Choose notes that are relevant to the user's query by meaning,
including synonyms and paraphrases. Do not invent notes.
Format:
{
  "matches": [
    {"index": 1, "score": 95},
    {"index": 3, "score": 70}
  ]
}
Use score 0-100. Return at most 5 matches.
"""
    payload = {
        "model": OLLAMA_MODEL,
        "format": "json",
        "stream": False,
        "messages": [
            {"role": "system", "content": system_prompt},
            {
                "role": "user",
                "content": json.dumps(
                    {"query": query, "notes": numbered},
                    ensure_ascii=False,
                ),
            },
        ],
        "options": {"temperature": 0.0, "num_ctx": 4096},
    }

    async with httpx.AsyncClient(timeout=45) as client:
        response = await client.post(f"{OLLAMA_URL}/api/chat", json=payload)
        response.raise_for_status()
        content = response.json()["message"]["content"]

    data = json.loads(content)
    matches = data.get("matches", [])
    ranked = []
    for match in matches:
        try:
            index = int(match.get("index")) - 1
            score = int(match.get("score", 0))
        except Exception:
            continue
        if 0 <= index < len(notes) and score >= 45:
            note = dict(notes[index])
            note["score"] = max(int(note.get("score", 0)), score)
            ranked.append(note)

    if not ranked:
        return notes[:limit]
    ranked.sort(key=lambda note: (int(note.get("score", 0)), note.get("created_at", "")), reverse=True)
    return ranked[:limit]


async def search_notes_contextual(
    chat_id: int,
    query: str,
    user_id: int | None = None,
    limit: int = 5,
) -> list[dict[str, Any]]:
    lexical = search_notes(chat_id, query, user_id=user_id, limit=NOTES_LLM_SEARCH_LIMIT)
    all_notes = list(reversed(iter_notes(chat_id, user_id=user_id)))[:NOTES_LLM_SEARCH_LIMIT]

    seen = set()
    candidates = []
    for note in [*lexical, *all_notes]:
        note_id = note.get("id")
        if note_id in seen:
            continue
        seen.add(note_id)
        candidates.append(note)

    if not candidates:
        return []

    try:
        return await rerank_notes_with_ollama(query, candidates, limit)
    except Exception:
        logger.warning("Ollama note rerank failed; using lexical search", exc_info=True)
        return lexical[:limit] if lexical else candidates[:limit]


def supported_knowledge_suffixes() -> set[str]:
    return {
        ".txt",
        ".md",
        ".markdown",
        ".csv",
        ".json",
        ".yaml",
        ".yml",
        ".log",
        ".html",
        ".htm",
        ".xml",
        ".pdf",
        ".docx",
    }


def extract_text_from_file(path: Path, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("Для PDF установи зависимость: pip install pypdf") from exc
        reader = PdfReader(str(path))
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        return "\n\n".join(pages).strip()

    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("Для DOCX установи зависимость: pip install python-docx") from exc
        document = Document(str(path))
        return "\n".join(paragraph.text for paragraph in document.paragraphs).strip()

    raw = path.read_bytes()
    for encoding in ("utf-8", "utf-8-sig", "cp1251", "latin-1"):
        try:
            return raw.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore").strip()


def chunk_text(text: str, max_chars: int = KNOWLEDGE_CHUNK_CHARS, overlap: int = KNOWLEDGE_CHUNK_OVERLAP) -> list[str]:
    import re

    clean = re.sub(r"\r\n?", "\n", text)
    clean = re.sub(r"\n{3,}", "\n\n", clean).strip()
    if not clean:
        return []

    paragraphs = [part.strip() for part in re.split(r"\n\s*\n", clean) if part.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if len(paragraph) > max_chars:
            if current:
                chunks.append(current.strip())
                current = ""
            start = 0
            while start < len(paragraph):
                chunks.append(paragraph[start : start + max_chars].strip())
                start += max(1, max_chars - overlap)
            continue

        candidate = f"{current}\n\n{paragraph}".strip() if current else paragraph
        if len(candidate) <= max_chars:
            current = candidate
        else:
            if current:
                chunks.append(current.strip())
            current = paragraph

    if current:
        chunks.append(current.strip())
    return chunks


def add_knowledge_document(
    chat_id: int,
    user_id: int | None,
    filename: str,
    text: str,
) -> dict[str, Any]:
    chunks = chunk_text(text)
    if not chunks:
        raise ValueError("В файле не нашел текста для базы знаний.")

    KNOWLEDGE_PATH.parent.mkdir(parents=True, exist_ok=True)
    document_id = f"kb-{int(datetime.now(ZoneInfo(TIMEZONE)).timestamp() * 1000)}"
    created_at = now_iso()
    with KNOWLEDGE_PATH.open("a", encoding="utf-8") as file:
        for index, chunk in enumerate(chunks, start=1):
            record = {
                "id": f"{document_id}:{index}",
                "document_id": document_id,
                "chat_id": chat_id,
                "user_id": user_id,
                "filename": filename,
                "chunk_index": index,
                "chunk_count": len(chunks),
                "text": chunk,
                "created_at": created_at,
            }
            file.write(
                json.dumps(
                    {"v": 1, "cipher": encrypt_text(json.dumps(record, ensure_ascii=False))},
                    ensure_ascii=False,
                )
                + "\n"
            )

    return {
        "document_id": document_id,
        "filename": filename,
        "chunk_count": len(chunks),
        "created_at": created_at,
    }


def iter_knowledge_chunks(chat_id: int, user_id: int | None = None) -> list[dict[str, Any]]:
    if not KNOWLEDGE_PATH.exists():
        return []

    chunks: list[dict[str, Any]] = []
    with KNOWLEDGE_PATH.open("r", encoding="utf-8") as file:
        for line in file:
            line = line.strip()
            if not line:
                continue
            try:
                wrapper = json.loads(line)
                chunk = json.loads(decrypt_text(str(wrapper["cipher"])))
            except Exception:
                logger.warning("Could not decrypt knowledge line", exc_info=True)
                continue
            if int(chunk.get("chat_id", 0)) != chat_id:
                continue
            if user_id is not None and chunk.get("user_id") not in (None, user_id):
                continue
            chunks.append(chunk)
    return chunks


def search_knowledge_chunks(chat_id: int, query: str, user_id: int | None = None, limit: int = 5) -> list[dict[str, Any]]:
    expanded_query = expand_note_query(query)
    ranked = []
    for chunk in iter_knowledge_chunks(chat_id, user_id=user_id):
        source_text = f"{chunk.get('filename', '')}\n{chunk.get('text', '')}"
        score = note_score(expanded_query, source_text)
        if score >= 20:
            ranked.append((score, chunk))
    ranked.sort(key=lambda item: (item[0], item[1].get("created_at", "")), reverse=True)
    return [chunk | {"score": score} for score, chunk in ranked[:limit]]


async def rerank_knowledge_with_ollama(query: str, chunks: list[dict[str, Any]], limit: int) -> list[dict[str, Any]]:
    if not KNOWLEDGE_LLM_SEARCH_ENABLED or not chunks:
        return chunks[:limit]

    numbered = []
    for index, chunk in enumerate(chunks, start=1):
        numbered.append(
            {
                "index": index,
                "filename": chunk.get("filename"),
                "chunk": str(chunk.get("text", ""))[:1200],
            }
        )

    system_prompt = """
You search a local encrypted knowledge base after chunks have been decrypted locally.
Return JSON only. Pick chunks relevant to the user's query by meaning.
Do not invent information. Prefer chunks that directly answer the query.
Format:
{"matches":[{"index":1,"score":95},{"index":2,"score":70}]}
Use score 0-100. Return at most 5 matches.
"""
    payload = {
        "model": OLLAMA_MODEL,
        "format": "json",
        "stream": False,
        "messages": [
            {"role": "system", "content": system_prompt},
            {
                "role": "user",
                "content": json.dumps({"query": query, "chunks": numbered}, ensure_ascii=False),
            },
        ],
        "options": {"temperature": 0.0, "num_ctx": 4096},
    }

    async with httpx.AsyncClient(timeout=60) as client:
        response = await client.post(f"{OLLAMA_URL}/api/chat", json=payload)
        response.raise_for_status()
        content = response.json()["message"]["content"]

    data = json.loads(content)
    ranked = []
    for match in data.get("matches", []):
        try:
            index = int(match.get("index")) - 1
            score = int(match.get("score", 0))
        except Exception:
            continue
        if 0 <= index < len(chunks) and score >= 40:
            chunk = dict(chunks[index])
            chunk["score"] = max(int(chunk.get("score", 0)), score)
            ranked.append(chunk)

    if not ranked:
        return chunks[:limit]
    ranked.sort(key=lambda chunk: (int(chunk.get("score", 0)), chunk.get("created_at", "")), reverse=True)
    return ranked[:limit]


async def search_knowledge_contextual(
    chat_id: int,
    query: str,
    user_id: int | None = None,
    limit: int = 5,
) -> list[dict[str, Any]]:
    lexical = search_knowledge_chunks(chat_id, query, user_id=user_id, limit=KNOWLEDGE_LLM_SEARCH_LIMIT)
    recent = list(reversed(iter_knowledge_chunks(chat_id, user_id=user_id)))[:KNOWLEDGE_LLM_SEARCH_LIMIT]

    seen = set()
    candidates = []
    for chunk in [*lexical, *recent]:
        chunk_id = chunk.get("id")
        if chunk_id in seen:
            continue
        seen.add(chunk_id)
        candidates.append(chunk)

    if not candidates:
        return []

    try:
        return await rerank_knowledge_with_ollama(query, candidates, limit)
    except Exception:
        logger.warning("Ollama knowledge rerank failed; using lexical search", exc_info=True)
        return lexical[:limit] if lexical else candidates[:limit]


def infer_knowledge_query(text: str) -> str | None:
    import re

    lowered = text.lower()
    if not any(marker in lowered for marker in ("баз", "знани", "файл", "документ")):
        return None
    if any(marker in lowered for marker in ("добав", "загру", "сохрани", "проиндекс")):
        return None

    patterns = (
        r"(?:найди|покажи|достань|поищи)\s+(?:в\s+)?(?:базе\s+знаний|файлах|документах|базе)\s*(?:про|о|об|по)?\s*(.+)",
        r"(?:что|какая|какие)\s+(?:в\s+)?(?:базе\s+знаний|файлах|документах|базе)\s*(?:про|о|об|по)?\s*(.+)",
    )
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE | re.DOTALL)
        if match:
            query = match.group(1).strip(" \n\t:;?.!-—")
            return query or None
    return None


def format_knowledge_results(query: str, chunks: list[dict[str, Any]]) -> str:
    if not chunks:
        return f"В базе знаний ничего не нашел по запросу: {query}"

    lines = [f"База знаний: {query}"]
    for index, chunk in enumerate(chunks, start=1):
        text = " ".join(str(chunk.get("text", "")).split())
        if len(text) > 700:
            text = text[:697].rstrip() + "..."
        lines.append(
            f"{index}. {chunk.get('filename', 'file')} "
            f"[{chunk.get('chunk_index')}/{chunk.get('chunk_count')}]"
        )
        lines.append(text)
    return "\n\n".join(lines)


def infer_note_create(text: str) -> str | None:
    import re

    lowered = text.lower()
    patterns = (
        r"(?:сделай|создай|запиши|добавь|сохрани)\s+заметк[ауи]\s*[:\-—,]?\s*(.+)",
        r"заметк[аи]\s*[:\-—,]\s*(.+)",
    )
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE | re.DOTALL)
        if match:
            note_text = match.group(1).strip(" \n\t:;-—")
            return note_text or None
    if lowered.strip() in {"сделай заметку", "создай заметку", "запиши заметку"}:
        return ""
    return None


def infer_note_query(text: str) -> str | None:
    import re

    lowered = text.lower()
    if not any(word in lowered for word in ("заметк", "записывал", "записано", "сохранял")):
        return None
    if any(word in lowered for word in ("сделай замет", "создай замет", "запиши замет", "добавь замет", "сохрани замет")):
        return None

    patterns = (
        r"(?:найди|покажи|достань|открой)\s+заметк\w*\s+(?:про|о|об|по)?\s*(.+)",
        r"(?:что|какие|какая)\s+(?:у меня\s+)?(?:есть\s+)?(?:в\s+)?заметк\w*\s+(?:про|о|об|по)?\s*(.+)",
        r"(?:что\s+)?(?:я\s+)?записывал\s+(?:про|о|об|по)?\s*(.+)",
    )
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE | re.DOTALL)
        if match:
            query = match.group(1).strip(" \n\t:;?.!-—")
            return query or None
    if "замет" in lowered:
        return lowered.replace("заметки", "").replace("заметку", "").replace("заметка", "").strip(" \n\t:;?.!-—") or None
    return None


def format_note_created(note: dict[str, Any]) -> str:
    return "\n".join(
        [
            "Заметка сохранена.",
            f"Когда: {format_msk_dt(note['created_at'])}",
            f"Текст: {note['text']}",
        ]
    )


def format_note_results(query: str, notes: list[dict[str, Any]]) -> str:
    if not notes:
        return f"Не нашел подходящих заметок по запросу: {query}"

    lines = [f"Нашел заметки по запросу: {query}"]
    for index, note in enumerate(notes, start=1):
        lines.append(
            f"{index}. {format_msk_dt(note['created_at'])} — {note['text']}"
        )
    return "\n".join(lines)


store = Store(DB_PATH)
scheduler = AsyncIOScheduler(timezone=ZoneInfo(TIMEZONE))
_whisper_model = None


def now_iso() -> str:
    return datetime.now(ZoneInfo(TIMEZONE)).isoformat(timespec="seconds")


def parse_dt(value: Any) -> datetime:
    if isinstance(value, dict):
        value = value.get("dateTime") or value.get("datetime") or value.get("date")
    if value is None:
        raise ValueError("missing datetime")
    if not isinstance(value, str):
        value = str(value)

    parsed = datetime.fromisoformat(value)
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=ZoneInfo(TIMEZONE))
    return parsed


def to_local_dt(value: Any) -> datetime:
    return parse_dt(value).astimezone(ZoneInfo(TIMEZONE))


def format_msk_dt(value: Any, *, include_date: bool = True) -> str:
    dt = to_local_dt(value)
    if include_date:
        return dt.strftime("%d.%m.%Y %H:%M МСК")
    return dt.strftime("%H:%M МСК")


def day_bounds(day: datetime) -> tuple[datetime, datetime]:
    start = day.astimezone(ZoneInfo(TIMEZONE)).replace(hour=0, minute=0, second=0, microsecond=0)
    return start, start + timedelta(days=1)


def has_negated_word(text: str, word: str) -> bool:
    import re

    lowered = text.lower()
    return bool(re.search(rf"\b(?:не|нет)\s+{re.escape(word)}\b", lowered))


def infer_explicit_date(text: str) -> datetime | None:
    import re

    lowered = text.lower()
    zone = ZoneInfo(TIMEZONE)
    today = datetime.now(zone).replace(hour=0, minute=0, second=0, microsecond=0)

    match = re.search(r"\b(\d{1,2})[.\-/](\d{1,2})(?:[.\-/](\d{2,4}))?\b", lowered)
    if match:
        day = int(match.group(1))
        month = int(match.group(2))
        year = int(match.group(3)) if match.group(3) else today.year
        if year < 100:
            year += 2000
        try:
            return datetime(year, month, day, tzinfo=zone)
        except ValueError:
            return None

    months = {
        "января": 1,
        "январь": 1,
        "февраля": 2,
        "февраль": 2,
        "марта": 3,
        "март": 3,
        "апреля": 4,
        "апрель": 4,
        "мая": 5,
        "май": 5,
        "июня": 6,
        "июнь": 6,
        "июля": 7,
        "июль": 7,
        "августа": 8,
        "август": 8,
        "сентября": 9,
        "сентябрь": 9,
        "октября": 10,
        "октябрь": 10,
        "ноября": 11,
        "ноябрь": 11,
        "декабря": 12,
        "декабрь": 12,
    }
    month_pattern = "|".join(months)
    match = re.search(rf"\b(\d{{1,2}})\s+({month_pattern})(?:\s+(\d{{2,4}}))?\b", lowered)
    if match:
        day = int(match.group(1))
        month = months[match.group(2)]
        year = int(match.group(3)) if match.group(3) else today.year
        if year < 100:
            year += 2000
        try:
            return datetime(year, month, day, tzinfo=zone)
        except ValueError:
            return None

    if "послезавтра" in lowered and not has_negated_word(lowered, "послезавтра"):
        return today + timedelta(days=2)
    if "завтра" in lowered and not has_negated_word(lowered, "завтра"):
        return today + timedelta(days=1)
    if "сегодня" in lowered and not has_negated_word(lowered, "сегодня"):
        return today

    return None


def reminder_range_from_text(text: str) -> tuple[str, datetime, datetime] | None:
    lowered = text.lower()
    reminder_words = ("напомин", "дел", "задач", "todo", "to-do")
    list_words = ("какие", "что", "покажи", "список", "есть", "будет", "будут", "заплан")
    if not any(word in lowered for word in reminder_words):
        return None
    if not any(word in lowered for word in list_words):
        return None

    zone = ZoneInfo(TIMEZONE)
    today_start = datetime.now(zone).replace(hour=0, minute=0, second=0, microsecond=0)

    start = infer_explicit_date(text)
    if start:
        start = start.replace(hour=0, minute=0, second=0, microsecond=0)
        if start.date() == today_start.date():
            return "Дела и напоминания на сегодня:", start, start + timedelta(days=1)
        if start.date() == (today_start + timedelta(days=1)).date():
            return "Дела и напоминания на завтра:", start, start + timedelta(days=1)
        if start.date() == (today_start + timedelta(days=2)).date():
            return "Дела и напоминания на послезавтра:", start, start + timedelta(days=1)
        return f"Дела и напоминания на {start.strftime('%d.%m.%Y')}:", start, start + timedelta(days=1)

    return "Дела и напоминания на сегодня:", today_start, today_start + timedelta(days=1)


def decode_unicode_escape_sequences(value: str) -> str:
    import re

    def replace_match(match: Any) -> str:
        return chr(int(match.group(1), 16))

    value = re.sub(r"\\u([0-9a-fA-F]{4})", replace_match, value)
    value = value.replace("\\n", "\n").replace("\\t", "\t")
    return value


def normalize_llm_strings(value: Any) -> Any:
    if isinstance(value, str):
        if "\\u" in value or "\\n" in value or "\\t" in value:
            return decode_unicode_escape_sequences(value)
        return value
    if isinstance(value, list):
        return [normalize_llm_strings(item) for item in value]
    if isinstance(value, dict):
        return {key: normalize_llm_strings(item) for key, item in value.items()}
    return value


def replace_date(value: Any, target_date: datetime) -> str:
    dt = parse_dt(value)
    corrected = dt.replace(
        year=target_date.year,
        month=target_date.month,
        day=target_date.day,
    )
    return corrected.isoformat()


def replace_time(value: Any, hour: int, minute: int) -> str:
    dt = parse_dt(value)
    corrected = dt.replace(hour=hour, minute=minute, second=0, microsecond=0)
    return corrected.isoformat()


def infer_target_date(text: str) -> datetime | None:
    lowered = text.lower()
    zone = ZoneInfo(TIMEZONE)
    today = datetime.now(zone).replace(hour=0, minute=0, second=0, microsecond=0)

    explicit = infer_explicit_date(text)
    if explicit:
        return explicit

    weekdays = {
        "понедельник": 0,
        "понедельника": 0,
        "вторник": 1,
        "вторника": 1,
        "среду": 2,
        "среда": 2,
        "четверг": 3,
        "четверга": 3,
        "пятницу": 4,
        "пятница": 4,
        "субботу": 5,
        "суббота": 5,
        "воскресенье": 6,
    }
    for word, weekday in weekdays.items():
        if word in lowered:
            days_ahead = (weekday - today.weekday()) % 7
            if days_ahead == 0:
                days_ahead = 7
            return today + timedelta(days=days_ahead)

    return None


def infer_weekday(text: str) -> int | None:
    lowered = text.lower()
    weekdays = {
        "понедельник": 0,
        "понедельникам": 0,
        "вторник": 1,
        "вторникам": 1,
        "среду": 2,
        "средам": 2,
        "четверг": 3,
        "четвергам": 3,
        "пятницу": 4,
        "пятницам": 4,
        "субботу": 5,
        "субботам": 5,
        "воскресенье": 6,
        "воскресеньям": 6,
    }
    for word, weekday in weekdays.items():
        if word in lowered:
            return weekday
    return None


def infer_start_time(text: str) -> tuple[int, int] | None:
    import re

    lowered = text.lower()
    matches = re.finditer(r"(?<!\d)(?:в|на|к)?\s*(\d{1,2})(?:[:.](\d{2}))?(?!\d)", lowered)
    for match in matches:
        prefix = lowered[max(0, match.start() - 8):match.start()]
        suffix = lowered[match.end():match.end() + 12]
        month_words = (
            "янв",
            "фев",
            "мар",
            "апр",
            "мая",
            "май",
            "июн",
            "июл",
            "авг",
            "сен",
            "окт",
            "ноя",
            "дек",
        )
        if (
            "за" in prefix
            or "течение" in prefix
            or "мин" in suffix
            or "час" in suffix
            or "дн" in suffix
            or "сут" in suffix
        ):
            continue
        if any(month in suffix for month in month_words):
            continue

        hour = int(match.group(1))
        minute = int(match.group(2) or "0")
        if 0 <= hour <= 23 and 0 <= minute <= 59:
            return hour, minute

    return None


def infer_recurrence_end(text: str, start: datetime) -> datetime:
    import re

    lowered = text.lower()
    duration_days = infer_duration_days(text)
    if duration_days:
        return start + timedelta(days=duration_days) - timedelta(seconds=1)

    if "до конца года" in lowered or "на весь год" in lowered:
        return start.replace(month=12, day=31, hour=23, minute=59, second=59)

    match = re.search(r"на\s+(\d{1,2})\s*(недел|недели|недель)", lowered)
    if match:
        return start + timedelta(weeks=int(match.group(1)))

    match = re.search(r"на\s+(\d{1,2})\s*(месяц|месяца|месяцев)", lowered)
    if match:
        return start + timedelta(days=31 * int(match.group(1)))

    return start + timedelta(days=365)


def infer_reminder_offset_minutes(text: str) -> int | None:
    import re

    lowered = text.lower()
    match = re.search(r"за\s+(\d{1,3})\s*(мин|минут)", lowered)
    if match:
        return int(match.group(1))

    match = re.search(r"за\s+(\d{1,2})\s*(час|часа|часов)", lowered)
    if match:
        return int(match.group(1)) * 60

    return None


def infer_duration_days(text: str) -> int | None:
    import re

    lowered = text.lower()
    match = re.search(r"(?:в\s+течение|на)\s+(\d{1,3})\s*(день|дня|дней|сутки|суток)", lowered)
    if match:
        return int(match.group(1))
    return None


def infer_hourly_countdown_reminders(text: str) -> dict[str, Any] | None:
    import re

    lowered = text.lower()
    if "напомин" not in lowered:
        return None
    if not re.search(r"кажд\w*\s+час", lowered):
        return None

    days = infer_duration_days(text)
    if not days:
        return None
    days = max(1, min(days, 60))

    zone = ZoneInfo(TIMEZONE)
    target_date = infer_target_date(text)
    if target_date is None:
        target_date = datetime.now(zone)
    start = target_date.astimezone(zone).replace(hour=9, minute=0, second=0, microsecond=0)
    target_time = infer_start_time(text)
    if target_time:
        start = start.replace(hour=target_time[0], minute=target_time[1])
    if start <= datetime.now(zone):
        start = datetime.now(zone).replace(minute=0, second=0, microsecond=0) + timedelta(hours=1)

    text_match = re.search(r"(?:с\s+текстом|текстом)\s+(.+?)(?:\s+начинай|\s+начиная|\s+в\s+течение|\s+каждый\s+день|$)", text, re.IGNORECASE)
    base_text = text_match.group(1).strip(" .,:;") if text_match else text.strip()
    if not base_text:
        base_text = text.strip()

    occurrences: list[str] = []
    texts: list[str] = []
    for day_index in range(days):
        countdown = days - day_index
        day_start = start + timedelta(days=day_index)
        for hour in range(24):
            remind_at = day_start + timedelta(hours=hour)
            occurrences.append(remind_at.isoformat())
            texts.append(f"{base_text}. Осталось дней: {countdown}")
            if len(occurrences) >= 500:
                break
        if len(occurrences) >= 500:
            break

    if not occurrences:
        return None

    return {
        "kind": "recurring_reminders",
        "reply": "Поставлю почасовые Telegram-напоминания с обратным отсчетом.",
        "recurring_reminders": {
            "text": base_text,
            "texts": texts,
            "occurrences": occurrences,
            "event_time": start.isoformat(),
            "until": parse_dt(occurrences[-1]).isoformat(),
            "count": len(occurrences),
        },
    }


def is_explicit_remind_at_time(text: str) -> bool:
    lowered = text.lower()
    return any(marker in lowered for marker in ("напомни в", "напомни на", "напомни к"))


def normalize_search_text(value: str) -> str:
    import re

    value = value.lower().replace("ё", "е")
    value = re.sub(r"[^a-zа-я0-9\s]", " ", value)
    return " ".join(value.split())


def expand_search_terms(query: str) -> set[str]:
    normalized = normalize_search_text(query)
    terms = {term for term in normalized.split() if len(term) >= 3}
    aliases = {
        "ван": {"ваня", "ваней", "ване", "иван", "иваном", "ивану"},
        "иван": {"ваня", "ваней", "ване", "иван", "иваном", "ивану"},
        "саша": {"саша", "сашей", "александр", "александром"},
        "александр": {"саша", "сашей", "александр", "александром"},
        "сергей": {"сергей", "сергеем", "сергею", "сережа", "сережей"},
        "дмитрий": {"дмитрий", "дмитрием", "диме", "дима", "димой"},
        "алексей": {"алексей", "алексеем", "леша", "лешей"},
        "екатерина": {"екатерина", "катя", "катей"},
    }

    expanded = set(terms)
    for term in list(terms):
        for key, values in aliases.items():
            if term.startswith(key) or key.startswith(term):
                expanded.update(values)

    stop_words = {
        "встреч",
        "созвон",
        "звонок",
        "перенеси",
        "перенести",
        "удали",
        "отмени",
        "событие",
        "календар",
    }
    return {term for term in expanded if not any(term.startswith(stop) for stop in stop_words)}


def event_search_blob(event: dict[str, Any]) -> str:
    parts = [
        event.get("summary", ""),
        event.get("description", ""),
        event.get("location", ""),
    ]
    for attendee in event.get("attendees", []) or []:
        parts.append(attendee.get("displayName", ""))
        parts.append(attendee.get("email", ""))
    return normalize_search_text(" ".join(parts))


def rank_calendar_events(events: list[dict[str, Any]], query: str | None) -> list[dict[str, Any]]:
    if not query:
        return events

    terms = expand_search_terms(query)
    if not terms:
        return events

    scored: list[tuple[int, dict[str, Any]]] = []
    for event in events:
        blob = event_search_blob(event)
        words = blob.split()
        score = 0
        for term in terms:
            if term in blob:
                score += 4
                continue
            if any(word.startswith(term) or term.startswith(word) for word in words):
                score += 2
                continue
            if fuzz:
                fuzzy_score = max((fuzz.ratio(term, word) for word in words), default=0)
                if fuzzy_score >= 86:
                    score += 2
                elif fuzzy_score >= 76:
                    score += 1
        if score > 0:
            scored.append((score, event))

    if not scored:
        return events

    scored.sort(key=lambda item: item[0], reverse=True)
    return [event for _score, event in scored]


def correct_event_datetime_from_text(intent: ParsedIntent, text: str) -> ParsedIntent:
    target_date = infer_target_date(text)
    target_time = infer_start_time(text)
    data = intent.data

    try:
        if intent.kind == "calendar_event":
            event = data["event"]
            old_start = parse_dt(event["start"])
            old_end = parse_dt(event["end"]) if event.get("end") else old_start + timedelta(hours=1)
            duration = old_end - old_start

            if target_date:
                event["start"] = replace_date(event["start"], target_date)
                if event.get("end"):
                    event["end"] = replace_date(event["end"], target_date)

            if target_time:
                event["start"] = replace_time(event["start"], *target_time)
                event["end"] = (parse_dt(event["start"]) + duration).isoformat()

            if event.get("reminder_minutes") is None:
                event["reminder_minutes"] = DEFAULT_EVENT_REMINDER_MINUTES

        elif intent.kind == "reminder":
            reminder = data["reminder"]
            if target_date:
                reminder["remind_at"] = replace_date(reminder["remind_at"], target_date)
            if target_time:
                reminder["remind_at"] = replace_time(reminder["remind_at"], *target_time)

        elif intent.kind == "reschedule_calendar_event":
            reschedule = data["reschedule_event"]
            if target_date and reschedule.get("new_start"):
                reschedule["new_start"] = replace_date(reschedule["new_start"], target_date)
                if reschedule.get("new_end"):
                    reschedule["new_end"] = replace_date(reschedule["new_end"], target_date)
            if target_time and reschedule.get("new_start"):
                reschedule["new_start"] = replace_time(reschedule["new_start"], *target_time)
    except Exception:
        logger.exception("Datetime correction failed")

    return intent


def infer_calendar_list_range(text: str) -> tuple[str, datetime, datetime] | None:
    lowered = text.lower()
    calendar_words = ("событ", "встреч", "календар", "план", "расписан")
    question_words = ("какие", "что", "покажи", "список", "есть", "будет", "будут")

    if not any(word in lowered for word in calendar_words):
        return None
    if not any(word in lowered for word in question_words):
        return None

    zone = ZoneInfo(TIMEZONE)
    today_start = datetime.now(zone).replace(hour=0, minute=0, second=0, microsecond=0)

    if "завтра" in lowered:
        start = today_start + timedelta(days=1)
        return "События на завтра:", start, start + timedelta(days=1)

    if "послезавтра" in lowered:
        start = today_start + timedelta(days=2)
        return "События на послезавтра:", start, start + timedelta(days=1)

    if "недел" in lowered or "7 дней" in lowered:
        return "События на ближайшие 7 дней:", today_start, today_start + timedelta(days=7)

    if "сегодня" in lowered:
        return "События на сегодня:", today_start, today_start + timedelta(days=1)

    if "ближай" in lowered:
        return "Ближайшие события:", datetime.now(zone), today_start + timedelta(days=14)

    return None


def is_capabilities_question(text: str) -> bool:
    lowered = text.lower()
    patterns = (
        "что ты умеешь",
        "что умеешь",
        "помощь",
        "help",
        "функции",
        "возможности",
        "как пользоваться",
    )
    return any(pattern in lowered for pattern in patterns)


def capabilities_text(section: str = "main") -> str:
    sections = {
        "main": "\n".join(
            [
                "Что я умею:",
                "1. Календарь: создавать, переносить, удалять и показывать события.",
                "2. Напоминания: отправлять Telegram-сообщение в нужное время.",
                "3. Голос: принимать голосовые и превращать их в задачи.",
                "4. Память: хранить локальный зашифрованный контекст.",
                "5. Действия: показывать последние выполненные операции.",
                "6. Заметки: сохранять короткие шифрованные заметки и искать их по смыслу.",
                "7. База знаний: добавлять файлы и искать ответы по их содержимому.",
            ]
        ),
        "calendar": "\n".join(
            [
                "Календарь:",
                "- создай встречу завтра в 15:00 с Иваном",
                "- перенеси встречу с Иваном на среду 10 утра",
                "- удали встречу с Иваном завтра",
                "- создай встречу каждый понедельник в 13:00 планерка",
                "- удали всю серию планерка",
                "- какие события есть на завтра",
                "Правило: 'создай встречу/событие' означает Google Calendar.",
                "Команды: /today, /tomorrow, /week",
            ]
        ),
        "reminders": "\n".join(
            [
                "Напоминания:",
                "- напомни завтра в 15:00 позвонить Ивану",
                "- напомни за 30 минут до встречи",
                "- каждый понедельник в 13:00 запросить отчет номер 1",
                "- поставь напоминание каждый час с текстом ... начиная с завтра в течение 10 дней",
                "- какие напоминания на сегодня",
                "- покажи задачи на послезавтра",
                "Правило: слово 'напомни' означает Telegram-сообщение, не событие календаря.",
                "Если сказано 'напомни' и отдельное время напоминания не указано, напомню за 15 минут до указанного времени.",
                "Повторяющиеся напоминания создаются на год вперед, если срок не указан.",
                "Почасовые серии с 'обратным отсчетом' создают отдельные напоминания с числом оставшихся дней.",
                "Команды: /reminders, /reminders завтра, /reminders послезавтра, /reminders 25.05.2026",
            ]
        ),
        "voice": "\n".join(
            [
                "Голос:",
                "- отправьте голосовое сообщение обычным Telegram voice",
                "- я распознаю текст локально через Whisper",
                "- дальше обработаю как текстовую задачу",
            ]
        ),
        "notes": "\n".join(
            [
                "Заметки:",
                "- сделай заметку клиент Иван предпочитает звонки после 14:00",
                "- найди заметку про Ивана",
                "- что я записывал про документы",
                "- /notes — последние заметки",
                "- /notes Иван — поиск по заметкам",
                "Заметки хранятся локально в data/notes.jsonl и шифруются.",
            ]
        ),
        "knowledge": "\n".join(
            [
                "База знаний:",
                "- пришли файл в Telegram, и я добавлю его в локальную базу знаний",
                "- поддержка: txt, md, csv, json, yaml, log, html, xml, pdf, docx",
                "- /kb — список загруженных файлов",
                "- /kb договор с Иваном — поиск по базе знаний",
                "- /knowledge регламент оплаты — то же самое",
                "- найди в базе знаний про условия договора",
                "- покажи в документах про отчет",
                "Как хранится: data/knowledge.jsonl, каждый фрагмент зашифрован локальным Fernet-ключом.",
                "Как ищет: быстрый локальный fuzzy-поиск + локальная Ollama-реранжировка, если включена.",
            ]
        ),
        "memory": "\n".join(
            [
                "Память и контроль:",
                "- /memory — статус локальной памяти",
                "- /forget_today — забыть сегодняшние сообщения",
                "- /debuglog — диагностика последних запросов",
                "- /stt — текущие настройки распознавания речи",
                "- /privacy — текущий режим приватности",
                "- /actions — последние действия",
                "- /last — последнее действие",
                "- /reminders [сегодня|завтра|послезавтра|25.05|25 мая] — дела и напоминания на дату",
                "История хранится локально и шифруется.",
            ]
        ),
    }
    return sections.get(section, sections["main"])


def capabilities_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        [
            [
                InlineKeyboardButton("Календарь", callback_data="help:calendar"),
                InlineKeyboardButton("Напоминания", callback_data="help:reminders"),
            ],
            [
                InlineKeyboardButton("Голос", callback_data="help:voice"),
                InlineKeyboardButton("Память", callback_data="help:memory"),
            ],
            [
                InlineKeyboardButton("Заметки", callback_data="help:notes"),
                InlineKeyboardButton("База знаний", callback_data="help:knowledge"),
            ],
        ]
    )


def infer_task_reminder(text: str) -> dict[str, Any] | None:
    lowered = text.lower()
    task_words = (
        "запросить",
        "попросить",
        "выслать",
        "отправить",
        "подготовить",
        "сделать",
        "проверить",
        "позвонить",
        "написать",
    )
    if not any(word in lowered for word in task_words):
        return None

    target_date = infer_target_date(text)
    target_time = infer_start_time(text)
    if not target_date or not target_time:
        return None

    remind_at = target_date.replace(
        hour=target_time[0],
        minute=target_time[1],
        second=0,
        microsecond=0,
    )
    if remind_at <= datetime.now(ZoneInfo(TIMEZONE)):
        return None

    return {
        "kind": "reminder",
        "reply": "Поставлю напоминание.",
        "reminder": {
            "text": text.strip(),
            "remind_at": remind_at.isoformat(),
        },
    }


def infer_direct_reminder(text: str) -> dict[str, Any] | None:
    lowered = text.lower()
    if "напомни" not in lowered:
        return None

    target_date = infer_target_date(text)
    target_time = infer_start_time(text)
    if not target_date or not target_time:
        return None

    event_time = target_date.replace(
        hour=target_time[0],
        minute=target_time[1],
        second=0,
        microsecond=0,
    )

    offset_minutes = infer_reminder_offset_minutes(text)
    if offset_minutes is not None:
        remind_at = event_time - timedelta(minutes=offset_minutes)
    elif is_explicit_remind_at_time(text):
        remind_at = event_time
    else:
        remind_at = event_time - timedelta(minutes=DEFAULT_EVENT_REMINDER_MINUTES)

    if remind_at <= datetime.now(ZoneInfo(TIMEZONE)):
        remind_at = event_time

    return {
        "kind": "reminder",
        "reply": "Поставлю Telegram-напоминание.",
        "reminder": {
            "text": text.strip(),
            "remind_at": remind_at.isoformat(),
        },
    }


def infer_recurring_task_reminders(text: str) -> dict[str, Any] | None:
    lowered = text.lower()
    recurrence_words = ("каждый", "каждую", "каждое", "каждые", "еженедельно")
    if not any(word in lowered for word in recurrence_words):
        return None

    weekday = infer_weekday(text)
    target_time = infer_start_time(text)
    if weekday is None or target_time is None:
        return None

    zone = ZoneInfo(TIMEZONE)
    now = datetime.now(zone)
    start_day = now.replace(hour=0, minute=0, second=0, microsecond=0)
    days_ahead = (weekday - start_day.weekday()) % 7
    first = start_day + timedelta(days=days_ahead)
    first = first.replace(hour=target_time[0], minute=target_time[1])
    if first <= now:
        first += timedelta(days=7)

    offset_minutes = infer_reminder_offset_minutes(text)
    end = infer_recurrence_end(text, first)

    occurrences = []
    current = first
    while current <= end and len(occurrences) < 60:
        remind_at = current
        if offset_minutes is not None:
            remind_at = current - timedelta(minutes=offset_minutes)
        occurrences.append(remind_at.isoformat())
        current += timedelta(days=7)

    if not occurrences:
        return None

    return {
        "kind": "recurring_reminders",
        "reply": "Поставлю повторяющиеся Telegram-напоминания.",
        "recurring_reminders": {
            "text": text.strip(),
            "occurrences": occurrences,
            "event_time": first.isoformat(),
            "until": end.isoformat(),
            "count": len(occurrences),
        },
    }


def extract_reminder_text(text: str) -> str:
    import re

    patterns = [
        r"(?:с\s+текстом|текстом)\s+(.+?)(?:\s+начинай|\s+начиная|\s+в\s+течение|\s+на\s+\d|\s+кажд|$)",
        r"(?:напомни\s+)?кажд\w*\s+(?:день|утро|вечер|ночь|час|понедельник|вторник|среду|четверг|пятницу|субботу|воскресенье)\s+(?:в|на|к)?\s*\d{1,2}(?:\s*час\w*)?(?::\d{2})?\s+(.+?)(?:\s+в\s+течение|\s+на\s+\d|$)",
        r"(?:напомни\s+)?кажд\w*\s+(?:утро|день|вечер|ночь)\s+(.+?)(?:\s+в\s+течение|\s+на\s+\d|$)",
        r"(?:напоминание|напомни)\s+(.+?)(?:\s+кажд|\s+начиная|\s+в\s+течение|\s+на\s+\d|$)",
    ]
    for pattern in patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            value = match.group(1).strip(" .,:;")
            if value:
                return value
    return text.strip()


def infer_general_recurring_reminders(text: str) -> dict[str, Any] | None:
    import re

    lowered = text.lower()
    if not any(word in lowered for word in ("напомин", "напомн", "присл", "сообщение", "запросить", "попросить", "позвонить", "написать", "сделать", "проверить")):
        return None
    if not any(word in lowered for word in ("каждый", "каждую", "каждое", "каждые", "ежедневно", "ежечасно", "еженедельно")):
        return None
    if "обратн" in lowered and "отсчет" in lowered:
        return None

    zone = ZoneInfo(TIMEZONE)
    now = datetime.now(zone)
    target_date = infer_target_date(text) or now
    start_day = target_date.astimezone(zone).replace(hour=0, minute=0, second=0, microsecond=0)
    target_time = infer_start_time(text)
    if target_time is None:
        if "утро" in lowered or "утром" in lowered:
            target_time = (9, 0)
        elif "день" in lowered or "днем" in lowered:
            target_time = (12, 0)
        elif "вечер" in lowered or "вечером" in lowered:
            target_time = (19, 0)
        elif "ночь" in lowered or "ночью" in lowered:
            target_time = (22, 0)

    interval = None
    step = None
    weekday = infer_weekday(text)

    every_hours = re.search(r"кажд\w*\s+(\d{1,2})\s*(?:час|часа|часов)", lowered)
    if every_hours:
        interval = "hours"
        step = int(every_hours.group(1))
    elif re.search(r"кажд\w*\s+час|ежечасно", lowered):
        interval = "hours"
        step = 1
    elif re.search(r"кажд\w*\s+(?:день|утро|вечер|ночь)|ежедневно", lowered):
        interval = "days"
        step = 1
    elif weekday is not None or "еженедельно" in lowered:
        interval = "weeks"
        step = 1
    else:
        return None

    if interval in {"days", "weeks"} and not target_time:
        return None

    if interval == "weeks":
        if weekday is None:
            weekday = start_day.weekday()
        days_ahead = (weekday - start_day.weekday()) % 7
        first = start_day + timedelta(days=days_ahead)
        first = first.replace(hour=target_time[0], minute=target_time[1])
        if first <= now:
            first += timedelta(days=7)
        delta = timedelta(weeks=step)
    elif interval == "days":
        first = start_day.replace(hour=target_time[0], minute=target_time[1])
        if first <= now:
            first += timedelta(days=1)
        delta = timedelta(days=step)
    else:
        first = start_day
        if target_time:
            first = first.replace(hour=target_time[0], minute=target_time[1])
        else:
            first = max(now, first).replace(minute=0, second=0, microsecond=0)
            if first <= now:
                first += timedelta(hours=1)
        delta = timedelta(hours=step)

    end = infer_recurrence_end(text, first)
    if interval == "hours" and not infer_duration_days(text) and not re.search(r"на\s+\d{1,3}\s*(час|часа|часов)", lowered):
        end = first + timedelta(days=1)

    occurrences = []
    current = first
    while current <= end and len(occurrences) < 500:
        occurrences.append(current.isoformat())
        current += delta

    if not occurrences:
        return None

    return {
        "kind": "recurring_reminders",
        "reply": "Поставлю повторяющиеся Telegram-напоминания.",
        "recurring_reminders": {
            "text": extract_reminder_text(text),
            "occurrences": occurrences,
            "event_time": first.isoformat(),
            "until": parse_dt(occurrences[-1]).isoformat(),
            "count": len(occurrences),
        },
    }


def infer_recurring_calendar_events(text: str) -> dict[str, Any] | None:
    lowered = text.lower()
    if not any(word in lowered for word in ("каждый", "каждую", "каждое", "каждые", "еженедельно")):
        return None
    if not any(word in lowered for word in ("создай", "добавь", "запланируй")):
        return None
    if not any(word in lowered for word in ("встреч", "событ", "созвон", "календар")):
        return None

    weekday = infer_weekday(text)
    target_time = infer_start_time(text)
    if weekday is None or target_time is None:
        return None

    zone = ZoneInfo(TIMEZONE)
    now = datetime.now(zone)
    start_day = now.replace(hour=0, minute=0, second=0, microsecond=0)
    days_ahead = (weekday - start_day.weekday()) % 7
    first = start_day + timedelta(days=days_ahead)
    first = first.replace(hour=target_time[0], minute=target_time[1])
    if first <= now:
        first += timedelta(days=7)

    end_limit = infer_recurrence_end(text, first)
    reminder_minutes = infer_reminder_offset_minutes(text)
    if reminder_minutes is None:
        reminder_minutes = DEFAULT_EVENT_REMINDER_MINUTES

    occurrences = []
    current = first
    while current <= end_limit and len(occurrences) < 60:
        occurrences.append(
            {
                "start": current.isoformat(),
                "end": (current + timedelta(hours=1)).isoformat(),
            }
        )
        current += timedelta(days=7)

    if not occurrences:
        return None

    title = text.strip()
    title = title.replace("создай", "").replace("добавь", "").replace("запланируй", "").strip()

    return {
        "kind": "recurring_calendar_events",
        "reply": "Создам повторяющиеся события в Google Calendar.",
        "recurring_calendar_events": {
            "title": title or text.strip(),
            "description": text.strip(),
            "reminder_minutes": reminder_minutes,
            "occurrences": occurrences,
            "until": end_limit.isoformat(),
            "count": len(occurrences),
        },
    }


def wants_delete_all(text: str) -> bool:
    lowered = text.lower()
    return any(
        marker in lowered
        for marker in (
            "удали все",
            "удалить все",
            "отмени все",
            "отменить все",
            "все такие",
            "всю серию",
            "серии",
        )
    )


def clean_reminder_delete_query(text: str) -> str:
    import re

    cleaned = text.lower()
    cleaned = re.sub(r"\b(удали|удалить|отмени|отменить|убери)\b", " ", cleaned)
    cleaned = re.sub(r"\b(все|всю|серия|серию|напоминание|напоминания|напоминаний|задачу|задачи|дело|дела|про|по)\b", " ", cleaned)
    cleaned = re.sub(r"\b(на|за|сегодня|завтра|послезавтра)\b", " ", cleaned)
    cleaned = re.sub(r"\b\d{1,2}[.\-/]\d{1,2}(?:[.\-/]\d{2,4})?\b", " ", cleaned)
    cleaned = re.sub(
        r"\b\d{1,2}\s+(января|январь|февраля|февраль|марта|март|апреля|апрель|мая|май|июня|июнь|июля|июль|августа|август|сентября|сентябрь|октября|октябрь|ноября|ноябрь|декабря|декабрь)(?:\s+\d{2,4})?\b",
        " ",
        cleaned,
    )
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" .,:;")
    return cleaned


def infer_delete_reminders(text: str) -> dict[str, Any] | None:
    lowered = text.lower()
    if not any(word in lowered for word in ("удали", "удалить", "отмени", "отменить", "убери")):
        return None
    if not any(word in lowered for word in ("напомин", "задач", "дел")):
        return None
    if any(word in lowered for word in ("календар", "встреч", "событ")):
        return None

    start = end = None
    target_date = infer_explicit_date(text)
    if target_date:
        start, end = day_bounds(target_date)

    delete_all = wants_delete_all(text) or any(word in lowered for word in ("серия", "серию"))
    query = clean_reminder_delete_query(text)
    return {
        "kind": "delete_reminders",
        "reply": "Удалю подходящие Telegram-напоминания.",
        "delete_reminders": {
            "query": query,
            "time_min": start.isoformat() if start else None,
            "time_max": end.isoformat() if end else None,
            "delete_all": delete_all,
        },
    }


def infer_delete_all_calendar_events(text: str) -> dict[str, Any] | None:
    lowered = text.lower()
    if not any(word in lowered for word in ("удали", "удалить", "отмени", "отменить")):
        return None
    if any(word in lowered for word in ("напомин", "задач", "дел")) and not any(
        word in lowered for word in ("календар", "встреч", "событ")
    ):
        return None
    if not wants_delete_all(text) and "кажд" not in lowered:
        return None

    zone = ZoneInfo(TIMEZONE)
    now = datetime.now(zone)
    time_min = now
    time_max = now + timedelta(days=365)

    target_date = infer_target_date(text)
    if target_date and "кажд" not in lowered and not wants_delete_all(text):
        time_min = target_date
        time_max = target_date + timedelta(days=1)

    return {
        "kind": "delete_calendar_event",
        "reply": "Удалю все найденные подходящие события.",
        "delete_event": {
            "query": text.strip(),
            "time_min": time_min.isoformat(),
            "time_max": time_max.isoformat(),
            "delete_all": True,
        },
    }


def is_allowed(update: Update) -> bool:
    if not ALLOWED_TELEGRAM_USER_IDS:
        return True
    user = update.effective_user
    return bool(user and user.id in ALLOWED_TELEGRAM_USER_IDS)


async def safe_send_typing(context: ContextTypes.DEFAULT_TYPE, chat_id: int) -> None:
    try:
        await context.bot.send_chat_action(chat_id, ChatAction.TYPING)
    except Exception:
        logger.debug("Could not send Telegram typing action", exc_info=True)


async def safe_send_message(app: Application, chat_id: int, text: str, **kwargs: Any) -> bool:
    for attempt in range(3):
        try:
            await app.bot.send_message(chat_id=chat_id, text=text, **kwargs)
            return True
        except Exception as exc:
            logger.warning(
                "Telegram send_message failed attempt %s/3: %s",
                attempt + 1,
                exc,
            )
            await asyncio.sleep(1 + attempt * 2)
    return False


async def safe_reply_text(message: Any, text: str, **kwargs: Any) -> bool:
    for attempt in range(3):
        try:
            await message.reply_text(text, **kwargs)
            return True
        except Exception as exc:
            logger.warning(
                "Telegram reply_text failed attempt %s/3: %s",
                attempt + 1,
                exc,
            )
            await asyncio.sleep(1 + attempt * 2)
    return False


def google_token_path_for_user(user_id: int | None) -> Path:
    if user_id is None:
        return TOKEN_PATH
    return GOOGLE_TOKENS_DIR / f"{user_id}.json"


def get_calendar_service(user_id: int | None = None) -> Any:
    from google.auth.transport.requests import Request
    from google.oauth2.credentials import Credentials
    from google.auth.exceptions import RefreshError
    from google_auth_oauthlib.flow import InstalledAppFlow
    from googleapiclient.discovery import build

    token_path = google_token_path_for_user(user_id)
    creds = None
    if token_path.exists():
        creds = Credentials.from_authorized_user_file(str(token_path), GOOGLE_SCOPES)
    elif user_id is not None and TOKEN_PATH.exists():
        # Backward-compatible migration for the first configured user.
        token_path.parent.mkdir(parents=True, exist_ok=True)
        token_path.write_text(TOKEN_PATH.read_text(encoding="utf-8"), encoding="utf-8")
        creds = Credentials.from_authorized_user_file(str(token_path), GOOGLE_SCOPES)

    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            try:
                creds.refresh(Request())
            except RefreshError:
                if token_path.exists():
                    token_path.unlink()
                creds = None
                raise RuntimeError(
                    f"Google Calendar token expired or was revoked for user {user_id}. "
                    f"Run `python google_calendar_setup.py --user-id {user_id}` to authorize again."
                )
        else:
            if not CREDENTIALS_PATH.exists():
                raise FileNotFoundError(
                    f"Put Google OAuth credentials into {CREDENTIALS_PATH}"
                )
            flow = InstalledAppFlow.from_client_secrets_file(
                str(CREDENTIALS_PATH),
                GOOGLE_SCOPES,
            )
            creds = flow.run_local_server(port=0)

        token_path.parent.mkdir(parents=True, exist_ok=True)
        token_path.write_text(creds.to_json(), encoding="utf-8")

    return build("calendar", "v3", credentials=creds)


async def google_api_to_thread(fn: Any, *, attempts: int = 3) -> Any:
    last_exc: Exception | None = None
    for attempt in range(attempts):
        try:
            return await asyncio.to_thread(fn)
        except Exception as exc:
            last_exc = exc
            logger.warning(
                "Google API call failed attempt %s/%s: %s",
                attempt + 1,
                attempts,
                exc,
            )
            await asyncio.sleep(1 + attempt * 2)
    assert last_exc is not None
    raise last_exc


async def create_calendar_event(action: dict[str, Any], user_id: int | None = None) -> str:
    event = action["event"]
    start = parse_dt(event["start"])
    end_raw = event.get("end")
    end = parse_dt(end_raw) if end_raw else start + timedelta(hours=1)
    reminder_minutes = event.get("reminder_minutes")
    if reminder_minutes is None:
        reminder_minutes = DEFAULT_EVENT_REMINDER_MINUTES

    body = {
        "summary": event["title"],
        "description": event.get("description", ""),
        "start": {"dateTime": start.isoformat(), "timeZone": TIMEZONE},
        "end": {"dateTime": end.isoformat(), "timeZone": TIMEZONE},
    }

    if event.get("location"):
        body["location"] = event["location"]

    body["reminders"] = {
        "useDefault": False,
        "overrides": [
            {"method": "popup", "minutes": int(reminder_minutes)},
        ],
    }

    service = await asyncio.to_thread(get_calendar_service, user_id)
    calendar_id = store.get_user_calendar_id(user_id) if user_id is not None else GOOGLE_CALENDAR_ID
    created = await google_api_to_thread(
        lambda: service.events()
        .insert(calendarId=calendar_id, body=body)
        .execute()
    )
    return created.get("htmlLink", "Событие создано.")


async def create_calendar_event_from_fields(
    title: str,
    start: str,
    end: str,
    description: str = "",
    location: str | None = None,
    reminder_minutes: int | None = None,
    user_id: int | None = None,
) -> str:
    action = {
        "event": {
            "title": title,
            "start": start,
            "end": end,
            "description": description,
            "location": location,
            "reminder_minutes": (
                reminder_minutes
                if reminder_minutes is not None
                else DEFAULT_EVENT_REMINDER_MINUTES
            ),
        }
    }
    return await create_calendar_event(action, user_id=user_id)


def calendar_event_summary(action: dict[str, Any], link: str | None = None) -> str:
    event = action["event"]
    start = parse_dt(event["start"])
    end_raw = event.get("end")
    end = parse_dt(end_raw) if end_raw else start + timedelta(hours=1)
    reminder_minutes = event.get("reminder_minutes")
    if reminder_minutes is None:
        reminder_minutes = DEFAULT_EVENT_REMINDER_MINUTES

    lines = [
        "Событие создано.",
        f"Название: {event['title']}",
        f"Когда: {start.strftime('%d.%m.%Y %H:%M')} - {end.strftime('%H:%M')}",
        f"Напоминание: за {reminder_minutes} мин.",
    ]

    if event.get("location"):
        lines.append(f"Место: {event['location']}")
    if event.get("description"):
        lines.append(f"Подробности: {event['description']}")
    if link:
        lines.append(f"Ссылка: {link}")

    return "\n".join(lines)


async def find_calendar_events(action: dict[str, Any], user_id: int | None = None) -> list[dict[str, Any]]:
    return await find_calendar_events_limited(action, limit=10, user_id=user_id)


async def find_calendar_events_limited(
    action: dict[str, Any],
    limit: int,
    user_id: int | None = None,
) -> list[dict[str, Any]]:
    query = action["delete_event"]
    zone = ZoneInfo(TIMEZONE)
    now = datetime.now(zone)
    time_min = parse_dt(query.get("time_min")) if query.get("time_min") else now
    time_max = (
        parse_dt(query.get("time_max"))
        if query.get("time_max")
        else now + timedelta(days=30)
    )

    events = await list_calendar_events(time_min, time_max, max_results=50, user_id=user_id)
    ranked = rank_calendar_events(events, query.get("query"))
    if not ranked and (query.get("time_min") or query.get("time_max")):
        events = await list_calendar_events(now, now + timedelta(days=30), max_results=100, user_id=user_id)
        ranked = rank_calendar_events(events, query.get("query"))
    return ranked[:limit]


async def list_calendar_events(
    time_min: datetime,
    time_max: datetime,
    max_results: int = 20,
    user_id: int | None = None,
) -> list[dict[str, Any]]:
    service = await asyncio.to_thread(get_calendar_service, user_id)
    calendar_id = store.get_user_calendar_id(user_id) if user_id is not None else GOOGLE_CALENDAR_ID

    def request() -> dict[str, Any]:
        return (
            service.events()
            .list(
                calendarId=calendar_id,
                timeMin=time_min.isoformat(),
                timeMax=time_max.isoformat(),
                singleEvents=True,
                orderBy="startTime",
                maxResults=max_results,
            )
            .execute()
        )

    result = await google_api_to_thread(request)
    return list(result.get("items", []))


async def delete_calendar_event_by_id(event_id: str, user_id: int | None = None) -> None:
    service = await asyncio.to_thread(get_calendar_service, user_id)
    calendar_id = store.get_user_calendar_id(user_id) if user_id is not None else GOOGLE_CALENDAR_ID
    await google_api_to_thread(
        lambda: service.events()
        .delete(calendarId=calendar_id, eventId=event_id)
        .execute()
    )


async def update_calendar_event_time(
    event_id: str,
    start: datetime,
    end: datetime,
    user_id: int | None = None,
) -> dict[str, Any]:
    service = await asyncio.to_thread(get_calendar_service, user_id)
    calendar_id = store.get_user_calendar_id(user_id) if user_id is not None else GOOGLE_CALENDAR_ID

    def request() -> dict[str, Any]:
        event = (
            service.events()
            .get(calendarId=calendar_id, eventId=event_id)
            .execute()
        )
        event["start"] = {"dateTime": start.isoformat(), "timeZone": TIMEZONE}
        event["end"] = {"dateTime": end.isoformat(), "timeZone": TIMEZONE}
        return (
            service.events()
            .update(calendarId=calendar_id, eventId=event_id, body=event)
            .execute()
        )

    return await google_api_to_thread(request)


def calendar_event_label(event: dict[str, Any]) -> str:
    start = event.get("start", {})
    start_value = start.get("dateTime") or start.get("date") or "без даты"
    return f"{event.get('summary', 'Без названия')} — {start_value}"


def format_calendar_events(events: list[dict[str, Any]], title: str) -> str:
    if not events:
        return f"{title}\nСобытий нет."

    lines = [title]
    for index, event in enumerate(events, start=1):
        start = event.get("start", {})
        end = event.get("end", {})
        start_raw = start.get("dateTime") or start.get("date")
        end_raw = end.get("dateTime") or end.get("date")

        if start_raw:
            start_dt = parse_dt(start_raw)
            if end_raw and "T" in end_raw:
                end_dt = parse_dt(end_raw)
                time_text = f"{start_dt.strftime('%d.%m %H:%M')}-{end_dt.strftime('%H:%M')}"
            elif "T" in start_raw:
                time_text = start_dt.strftime("%d.%m %H:%M")
            else:
                time_text = start_dt.strftime("%d.%m")
        else:
            time_text = "без даты"

        summary = event.get("summary", "Без названия")
        location = event.get("location")
        line = f"{index}. {time_text} — {summary}"
        if location:
            line += f" ({location})"
        lines.append(line)

    return "\n".join(lines)


def build_memory_context(chat_id: int) -> str:
    if not MEMORY_ENABLED:
        return "Memory is disabled."

    message_rows = list(reversed(store.recent_memory_messages(chat_id, MEMORY_RECENT_MESSAGES)))
    action_rows = list(reversed(store.recent_action_history(chat_id, MEMORY_RECENT_ACTIONS)))
    lines: list[str] = []

    if message_rows:
        lines.append("Recent encrypted local chat memory, decrypted for this prompt:")
        for row in message_rows:
            try:
                text = decrypt_text(row["text_cipher"])
            except Exception:
                text = "[could not decrypt]"
            lines.append(f"- {row['created_at']} {row['role']}: {text[:500]}")

    if action_rows:
        lines.append("Recent local action history:")
        for row in action_rows:
            try:
                summary = decrypt_text(row["summary_cipher"])
            except Exception:
                summary = "[could not decrypt]"
            lines.append(f"- {row['created_at']} {row['kind']}: {summary[:500]}")

    if not lines:
        return "No relevant local memory yet."

    return "\n".join(lines)


async def ask_ollama_for_intent(text: str, chat_id: int) -> ParsedIntent:
    memory_context = build_memory_context(chat_id)
    system_prompt = f"""
You are a local private assistant for a Telegram bot.
Convert the user's Russian or English message into one JSON object only.
Current datetime: {now_iso()}
Timezone: {TIMEZONE}

Local context:
{memory_context}

Supported JSON:
1. Calendar event:
{{
  "kind": "calendar_event",
  "reply": "short Russian confirmation summary",
  "event": {{
    "title": "string",
    "start": "ISO-8601 datetime with timezone",
    "end": "ISO-8601 datetime with timezone or null",
    "location": "string or null",
    "description": "string or null",
    "reminder_minutes": 15
  }}
}}

2. Telegram reminder:
{{
  "kind": "reminder",
  "reply": "short Russian confirmation summary",
  "reminder": {{
    "text": "what to send later",
    "remind_at": "ISO-8601 datetime with timezone"
  }}
}}

3. Delete/cancel calendar event:
{{
  "kind": "delete_calendar_event",
  "reply": "short Russian confirmation summary",
  "delete_event": {{
    "query": "event title, person, or keyword to search for",
    "time_min": "ISO-8601 datetime with timezone or null",
    "time_max": "ISO-8601 datetime with timezone or null",
    "delete_all": false
  }}
}}

4. Reschedule/move calendar event:
{{
  "kind": "reschedule_calendar_event",
  "reply": "short Russian confirmation summary",
  "reschedule_event": {{
    "query": "event title, person, or keyword to search for",
    "time_min": "ISO-8601 datetime with timezone or null",
    "time_max": "ISO-8601 datetime with timezone or null",
    "new_start": "ISO-8601 datetime with timezone",
    "new_end": "ISO-8601 datetime with timezone or null"
  }}
}}

5. Normal chat:
{{
  "kind": "chat",
  "reply": "short useful Russian answer"
}}

Rules:
- If the user says "напомни", classify it as "reminder", not "calendar_event".
- If date or time is missing for a new event/reminder, use kind "chat" and ask one clarifying question.
- For calendar_event, if the user does not specify reminder timing, set reminder_minutes to {DEFAULT_EVENT_REMINDER_MINUTES}.
- For delete_calendar_event, use a reasonable search range. If the user says tomorrow, set that day's range.
- If the user asks to delete all matching events or the whole series, set delete_all to true.
- If delete request has no useful title/time/person, use kind "chat" and ask what to delete.
- For reschedule_calendar_event, if the new date/time is missing, use kind "chat" and ask when to move it.
- For reschedule_calendar_event, use a reasonable search range to find the original event.
- Use local context to resolve references like "это", "последняя встреча", "вторая", "как выше".
- Output readable UTF-8 Russian text. Do not output literal unicode escape sequences like "\\u041f".
- Never invent card numbers, payment details, contacts, emails, or private data.
- Do not execute actions yourself. Only describe the parsed action.
- Return JSON only, no markdown.
""".strip()

    payload = {
        "model": OLLAMA_MODEL,
        "format": "json",
        "stream": False,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": text},
        ],
        "options": {
            "temperature": 0.1,
            "num_ctx": 2048,
        },
    }

    async with httpx.AsyncClient(timeout=120) as client:
        response = await client.post(f"{OLLAMA_URL}/api/chat", json=payload)
        response.raise_for_status()
        content = response.json()["message"]["content"]

    try:
        data = json.loads(content)
    except json.JSONDecodeError:
        logger.warning("Ollama returned non-JSON: %s", content)
        return ParsedIntent(kind="chat", data={}, reply=content.strip())

    data = normalize_llm_strings(data)
    kind = str(data.get("kind", "chat"))
    reply = str(data.get("reply", "")).strip() or "Понял."
    intent = ParsedIntent(kind=kind, data=data, reply=reply)
    return correct_event_datetime_from_text(intent, text)


def format_action(action: dict[str, Any]) -> str:
    if action["kind"] == "calendar_event":
        event = action["event"]
        lines = [
            "Создать событие в Google Calendar?",
            f"Название: {event['title']}",
            f"Начало: {event['start']}",
            f"Конец: {event.get('end') or 'через 1 час'}",
        ]
        if event.get("location"):
            lines.append(f"Место: {event['location']}")
        if event.get("reminder_minutes") is not None:
            lines.append(f"Напомнить за: {event['reminder_minutes']} мин.")
        return "\n".join(lines)

    if action["kind"] == "reminder":
        reminder = action["reminder"]
        return "\n".join(
            [
                "Поставить напоминание в Telegram?",
                f"Когда: {format_msk_dt(reminder['remind_at'])}",
                f"Текст: {reminder['text']}",
            ]
        )

    if action["kind"] == "recurring_reminders":
        recurring = action["recurring_reminders"]
        occurrences = recurring["occurrences"]
        return "\n".join(
            [
                "Поставить повторяющиеся Telegram-напоминания?",
                f"Количество: {recurring['count']}",
                f"Первое: {format_msk_dt(occurrences[0])}",
                f"Последнее: {format_msk_dt(occurrences[-1])}",
                f"Текст: {recurring['text']}",
            ]
        )

    if action["kind"] == "recurring_calendar_events":
        recurring = action["recurring_calendar_events"]
        occurrences = recurring["occurrences"]
        return "\n".join(
            [
                "Создать повторяющиеся события в Google Calendar?",
                f"Количество: {recurring['count']}",
                f"Первое: {format_msk_dt(occurrences[0]['start'])}",
                f"Последнее: {format_msk_dt(occurrences[-1]['start'])}",
                f"Напоминание: за {recurring['reminder_minutes']} мин.",
                f"Название: {recurring['title']}",
            ]
        )

    if action["kind"] == "delete_calendar_event":
        delete_event = action["delete_event"]
        return "\n".join(
            [
                "Найти и удалить события в Google Calendar?" if delete_event.get("delete_all") else "Найти и удалить событие в Google Calendar?",
                f"Запрос: {delete_event.get('query') or 'без ключевых слов'}",
                f"С начала: {delete_event.get('time_min') or 'сейчас'}",
                f"До: {delete_event.get('time_max') or 'через 30 дней'}",
                f"Режим: {'удалить все найденные' if delete_event.get('delete_all') else 'удалить одно событие'}",
            ]
        )

    if action["kind"] == "delete_reminders":
        delete_reminders = action["delete_reminders"]
        return "\n".join(
            [
                "Удалить Telegram-напоминания?",
                f"Запрос: {delete_reminders.get('query') or 'без ключевых слов'}",
                f"С начала: {format_msk_dt(delete_reminders['time_min']) if delete_reminders.get('time_min') else 'любая дата'}",
                f"До: {format_msk_dt(delete_reminders['time_max']) if delete_reminders.get('time_max') else 'любая дата'}",
                f"Режим: {'удалить все найденные' if delete_reminders.get('delete_all') else 'выбрать/удалить одно'}",
            ]
        )

    if action["kind"] == "reschedule_calendar_event":
        reschedule_event = action["reschedule_event"]
        return "\n".join(
            [
                "Найти и перенести событие в Google Calendar?",
                f"Запрос: {reschedule_event.get('query') or 'без ключевых слов'}",
                f"Искать с: {reschedule_event.get('time_min') or 'сейчас'}",
                f"Искать до: {reschedule_event.get('time_max') or 'через 30 дней'}",
                f"Новое начало: {format_msk_dt(reschedule_event.get('new_start'))}",
                f"Новый конец: {format_msk_dt(reschedule_event['new_end']) if reschedule_event.get('new_end') else 'сохранить длительность'}",
            ]
        )

    return action.get("reply", "Понял.")


async def schedule_reminder(app: Application, reminder_id: int, chat_id: int, text: str, remind_at: str) -> None:
    run_date = parse_dt(remind_at)
    scheduler.add_job(
        send_reminder,
        "date",
        id=f"reminder:{reminder_id}",
        replace_existing=True,
        run_date=run_date,
        args=[app, reminder_id, chat_id, text],
    )


async def send_reminder(app: Application, reminder_id: int, chat_id: int, text: str) -> None:
    sent = await safe_send_message(app, chat_id=chat_id, text=f"Напоминание: {text}")
    if sent:
        store.mark_reminder_sent(reminder_id)
        export_reminders_markdown(chat_id)
    else:
        logger.error("Reminder %s was not sent after retries", reminder_id)


def selected_stt_providers() -> list[str]:
    if PRIVACY_MODE == "strict":
        return ["local"]

    providers: list[str] = []
    for provider in (STT_PROVIDER, STT_FALLBACK_PROVIDER, "local"):
        provider = provider.strip().lower()
        if provider and provider not in providers:
            providers.append(provider)
    return providers


async def transcribe_with_openai_compatible(
    ogg_path: Path,
    *,
    api_key: str,
    url: str,
    model: str,
    provider_name: str,
) -> str:
    if not api_key:
        raise RuntimeError(f"{provider_name} API key is not configured")

    headers = {"Authorization": f"Bearer {api_key}"}
    data = {
        "model": model,
        "language": "ru",
        "response_format": "json",
    }

    async with httpx.AsyncClient(timeout=90) as client:
        with ogg_path.open("rb") as audio_file:
            content_type = mimetypes.guess_type(str(ogg_path))[0] or "application/octet-stream"
            files = {"file": (ogg_path.name, audio_file, content_type)}
            response = await client.post(url, headers=headers, data=data, files=files)
            response.raise_for_status()

    payload = response.json()
    text = str(payload.get("text", "")).strip()
    if not text:
        raise RuntimeError(f"{provider_name} returned empty transcript")
    return text


async def transcribe_with_local_whisper(ogg_path: Path) -> str:
    global _whisper_model

    if _whisper_model is None:
        from faster_whisper import WhisperModel

        _whisper_model = await asyncio.to_thread(
            WhisperModel,
            WHISPER_MODEL,
            device="cpu",
            compute_type="int8",
        )

    def run_transcription() -> str:
        segments, _info = _whisper_model.transcribe(
            str(ogg_path),
            language="ru",
            vad_filter=True,
        )
        return " ".join(segment.text.strip() for segment in segments).strip()

    return await asyncio.to_thread(run_transcription)


async def transcribe_file(ogg_path: Path) -> tuple[str, str]:
    errors: list[str] = []
    for provider in selected_stt_providers():
        try:
            if provider == "groq":
                text = await transcribe_with_openai_compatible(
                    ogg_path,
                    api_key=GROQ_API_KEY,
                    url=GROQ_STT_URL,
                    model=GROQ_STT_MODEL,
                    provider_name="Groq",
                )
            elif provider == "openai":
                text = await transcribe_with_openai_compatible(
                    ogg_path,
                    api_key=OPENAI_API_KEY,
                    url=OPENAI_STT_URL,
                    model=OPENAI_STT_MODEL,
                    provider_name="OpenAI",
                )
            elif provider == "local":
                text = await transcribe_with_local_whisper(ogg_path)
            else:
                raise RuntimeError(f"Unknown STT provider: {provider}")

            return text, provider
        except ImportError:
            raise
        except Exception as exc:
            logger.warning("STT provider %s failed: %s", provider, exc)
            errors.append(f"{provider}: {exc}")

    raise RuntimeError("; ".join(errors) or "No STT providers configured")


async def transcribe_voice(update: Update, context: ContextTypes.DEFAULT_TYPE) -> tuple[str, str]:

    if not update.message or not update.message.voice:
        raise ValueError("No voice message")
    if update.message.voice.duration and update.message.voice.duration > STT_MAX_VOICE_SECONDS:
        raise ValueError(
            f"Голосовое слишком длинное: {update.message.voice.duration} сек. Лимит: {STT_MAX_VOICE_SECONDS} сек."
        )

    await safe_send_typing(context, update.effective_chat.id)
    voice_file = await context.bot.get_file(update.message.voice.file_id)

    with tempfile.TemporaryDirectory() as temp_dir:
        ogg_path = Path(temp_dir) / "voice.ogg"
        await voice_file.download_to_drive(custom_path=str(ogg_path))
        return await transcribe_file(ogg_path)


async def maybe_handle_notes(
    text: str,
    *,
    chat_id: int,
    user_id: int | None,
    reply_target: Any,
    debug_id: int | None,
) -> bool:
    note_text = infer_note_create(text)
    if note_text is not None:
        if not note_text:
            reply = "Что записать в заметку? Скажи: сделай заметку <текст>."
            store.update_debug_log(debug_id, intent_kind="note_create", result="missing_text")
            await safe_reply_text(reply_target, reply)
            store.add_memory_message(chat_id=chat_id, user_id=None, role="assistant", text=reply)
            return True

        note = add_note(chat_id, user_id, note_text)
        reply = format_note_created(note)
        store.update_debug_log(
            debug_id,
            intent_kind="note_create",
            intent_payload={"text": note_text},
            result="created",
        )
        store.add_action_history(chat_id, "note_create", reply, {"note_id": note["id"]})
        await safe_reply_text(reply_target, reply)
        store.add_memory_message(chat_id=chat_id, user_id=None, role="assistant", text=reply)
        return True

    note_query = infer_note_query(text)
    if note_query:
        notes = await search_notes_contextual(chat_id, note_query, user_id=user_id)
        reply = format_note_results(note_query, notes)
        store.update_debug_log(
            debug_id,
            intent_kind="note_search",
            intent_payload={"query": note_query, "count": len(notes)},
            result="replied",
        )
        await safe_reply_text(reply_target, reply)
        store.add_memory_message(chat_id=chat_id, user_id=None, role="assistant", text=reply)
        return True

    return False


async def maybe_handle_knowledge_query(
    text: str,
    *,
    chat_id: int,
    user_id: int | None,
    reply_target: Any,
    debug_id: int | None,
) -> bool:
    query = infer_knowledge_query(text)
    if not query:
        return False

    chunks = await search_knowledge_contextual(chat_id, query, user_id=user_id)
    reply = format_knowledge_results(query, chunks)
    store.update_debug_log(
        debug_id,
        intent_kind="knowledge_search",
        intent_payload={"query": query, "count": len(chunks)},
        result="replied",
    )
    await safe_reply_text(reply_target, reply)
    store.add_memory_message(chat_id=chat_id, user_id=None, role="assistant", text=reply)
    return True


async def handle_text(
    text: str,
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
    *,
    source: str = "text",
    stt_provider: str | None = None,
    debug_id: int | None = None,
) -> None:
    if not update.message or not update.effective_user or not update.effective_chat:
        return
    if debug_id is None:
        debug_id = store.add_debug_log(
            chat_id=update.effective_chat.id,
            user_id=update.effective_user.id,
            source=source,
            stt_provider=stt_provider,
            input_text=text,
        )

    store.add_memory_message(
        chat_id=update.effective_chat.id,
        user_id=update.effective_user.id,
        role="user",
        text=text,
    )
    await safe_send_typing(context, update.effective_chat.id)

    if is_capabilities_question(text):
        reply = capabilities_text()
        store.update_debug_log(debug_id, intent_kind="capabilities", result="replied")
        await safe_reply_text(update.message, reply, reply_markup=capabilities_keyboard())
        store.add_memory_message(
            chat_id=update.effective_chat.id,
            user_id=None,
            role="assistant",
            text=reply,
        )
        return

    if await maybe_handle_notes(
        text,
        chat_id=update.effective_chat.id,
        user_id=update.effective_user.id,
        reply_target=update.message,
        debug_id=debug_id,
    ):
        return

    if await maybe_handle_knowledge_query(
        text,
        chat_id=update.effective_chat.id,
        user_id=update.effective_user.id,
        reply_target=update.message,
        debug_id=debug_id,
    ):
        return

    reminder_range = reminder_range_from_text(text)
    if reminder_range:
        title, start, end = reminder_range
        export_reminders_markdown(update.effective_chat.id)
        reply = format_reminders_for_range(update.effective_chat.id, title, start, end)
        store.update_debug_log(debug_id, intent_kind="reminder_list", result="replied")
        await safe_reply_text(update.message, reply)
        store.add_memory_message(
            chat_id=update.effective_chat.id,
            user_id=None,
            role="assistant",
            text=reply,
        )
        return

    calendar_range = infer_calendar_list_range(text)
    if calendar_range:
        title, start, end = calendar_range
        try:
            events = await list_calendar_events(start, end, user_id=update.effective_user.id)
            reply = format_calendar_events(events, title)
            store.update_debug_log(debug_id, intent_kind="calendar_list", result="replied")
        except Exception as exc:
            logger.exception("Calendar listing failed")
            reply = f"Не получилось прочитать календарь: {exc}"
            store.update_debug_log(
                debug_id,
                intent_kind="calendar_list",
                result="error",
                error=str(exc),
            )

        await safe_reply_text(update.message, reply)
        store.add_memory_message(
            chat_id=update.effective_chat.id,
            user_id=None,
            role="assistant",
            text=reply,
        )
        return

    task_reminder = (
        infer_delete_reminders(text)
        or infer_delete_all_calendar_events(text)
        or infer_recurring_calendar_events(text)
        or infer_hourly_countdown_reminders(text)
        or infer_general_recurring_reminders(text)
        or infer_recurring_task_reminders(text)
        or infer_direct_reminder(text)
        or infer_task_reminder(text)
    )
    if task_reminder:
        intent = ParsedIntent(
            kind="reminder",
            data=task_reminder,
            reply=task_reminder["reply"],
        )
        store.update_debug_log(
            debug_id,
            intent_kind=intent.kind,
            intent_payload=intent.data,
            result="auto_confirm_execute" if AUTO_CONFIRM_ACTIONS else "pending_confirmation",
        )
        if AUTO_CONFIRM_ACTIONS:
            await execute_action(
                intent.data,
                chat_id=update.effective_chat.id,
                user_id=update.effective_user.id,
                app=context.application,
                reply_target=update.message,
            )
            return

        action_id = store.add_pending(
            chat_id=update.effective_chat.id,
            user_id=update.effective_user.id,
            action=intent.data,
        )
        reply = format_action(intent.data)
        keyboard = InlineKeyboardMarkup(
            [
                [
                    InlineKeyboardButton("Подтвердить", callback_data=f"confirm:{action_id}"),
                    InlineKeyboardButton("Отмена", callback_data=f"cancel:{action_id}"),
                ]
            ]
        )
        await safe_reply_text(update.message, reply, reply_markup=keyboard)
        store.add_memory_message(
            chat_id=update.effective_chat.id,
            user_id=None,
            role="assistant",
            text=reply,
        )
        return

    try:
        intent = await ask_ollama_for_intent(text, update.effective_chat.id)
    except httpx.RequestError:
        store.update_debug_log(debug_id, result="error", error="ollama_request_error")
        await safe_reply_text(update.message, "Не могу подключиться к Ollama. Проверьте, что ollama serve запущен.")
        return
    except Exception as exc:
        logger.exception("Intent parsing failed")
        store.update_debug_log(debug_id, result="error", error=str(exc))
        await safe_reply_text(update.message, "Не смог разобрать запрос. Попробуйте сказать проще и с датой/временем.")
        return

    if intent.kind in {
        "calendar_event",
        "reminder",
        "recurring_reminders",
        "recurring_calendar_events",
        "delete_calendar_event",
        "delete_reminders",
        "reschedule_calendar_event",
    }:
        action_summary = format_action(intent.data)
        store.update_debug_log(
            debug_id,
            intent_kind=intent.kind,
            intent_payload=intent.data,
            result="auto_confirm_execute" if AUTO_CONFIRM_ACTIONS else "pending_confirmation",
        )
        store.add_action_history(
            update.effective_chat.id,
            intent.kind,
            action_summary,
            intent.data,
        )
        if AUTO_CONFIRM_ACTIONS:
            await execute_action(
                intent.data,
                chat_id=update.effective_chat.id,
                user_id=update.effective_user.id,
                app=context.application,
                reply_target=update.message,
            )
            return

        action_id = store.add_pending(
            chat_id=update.effective_chat.id,
            user_id=update.effective_user.id,
            action=intent.data,
        )
        keyboard = InlineKeyboardMarkup(
            [
                [
                    InlineKeyboardButton("Подтвердить", callback_data=f"confirm:{action_id}"),
                    InlineKeyboardButton("Отмена", callback_data=f"cancel:{action_id}"),
                ]
            ]
        )
        reply = action_summary
        await safe_reply_text(update.message, reply, reply_markup=keyboard)
        store.add_memory_message(
            chat_id=update.effective_chat.id,
            user_id=None,
            role="assistant",
            text=reply,
        )
        return

    store.update_debug_log(
        debug_id,
        intent_kind=intent.kind,
        intent_payload=intent.data,
        result="chat_reply",
    )
    await safe_reply_text(update.message, intent.reply)
    store.add_memory_message(
        chat_id=update.effective_chat.id,
        user_id=None,
        role="assistant",
        text=intent.reply,
    )


async def execute_action(
    parsed: dict[str, Any],
    chat_id: int,
    user_id: int | None,
    app: Application,
    reply_target: Any,
) -> None:
    try:
        if parsed["kind"] == "calendar_event":
            link = await create_calendar_event(parsed, user_id=user_id)
            await safe_reply_text(reply_target, calendar_event_summary(parsed, link))
        elif parsed["kind"] == "reminder":
            reminder = parsed["reminder"]
            reminder_id = store.add_reminder(
                chat_id=chat_id,
                text=reminder["text"],
                remind_at=reminder["remind_at"],
            )
            await schedule_reminder(
                app,
                reminder_id,
                chat_id,
                reminder["text"],
                reminder["remind_at"],
            )
            export_reminders_markdown(chat_id)
            await safe_reply_text(
                reply_target,
                "\n".join(
                    [
                        "Напоминание поставлено.",
                        f"Когда: {format_msk_dt(reminder['remind_at'])}",
                        f"Текст: {reminder['text']}",
                    ]
                )
            )
        elif parsed["kind"] == "recurring_reminders":
            recurring = parsed["recurring_reminders"]
            reminder_texts = recurring.get("texts") or []
            for index, remind_at in enumerate(recurring["occurrences"]):
                reminder_text = reminder_texts[index] if index < len(reminder_texts) else recurring["text"]
                reminder_id = store.add_reminder(
                    chat_id=chat_id,
                    text=reminder_text,
                    remind_at=remind_at,
                )
                await schedule_reminder(
                    app,
                    reminder_id,
                    chat_id,
                    reminder_text,
                    remind_at,
                )
            export_reminders_markdown(chat_id)
            await safe_reply_text(
                reply_target,
                "\n".join(
                    [
                        "Повторяющиеся напоминания поставлены.",
                        f"Количество: {recurring['count']}",
                        f"Первое: {format_msk_dt(recurring['occurrences'][0])}",
                        f"Последнее: {format_msk_dt(recurring['occurrences'][-1])}",
                    ]
                )
            )
        elif parsed["kind"] == "recurring_calendar_events":
            recurring = parsed["recurring_calendar_events"]
            links = []
            for occurrence in recurring["occurrences"]:
                link = await create_calendar_event_from_fields(
                    title=recurring["title"],
                    start=occurrence["start"],
                    end=occurrence["end"],
                    description=recurring.get("description", ""),
                    reminder_minutes=recurring.get("reminder_minutes"),
                    user_id=user_id,
                )
                links.append(link)
            await safe_reply_text(
                reply_target,
                "\n".join(
                    [
                        "Повторяющиеся события созданы.",
                        f"Количество: {recurring['count']}",
                        f"Первое: {format_msk_dt(recurring['occurrences'][0]['start'])}",
                        f"Последнее: {format_msk_dt(recurring['occurrences'][-1]['start'])}",
                        f"Напоминание: за {recurring['reminder_minutes']} мин.",
                    ]
                )
            )
        elif parsed["kind"] == "delete_calendar_event":
            if parsed["delete_event"].get("delete_all"):
                events = await find_calendar_events_limited(parsed, limit=100, user_id=user_id)
                if not events:
                    await safe_reply_text(reply_target, "Не нашел подходящих событий в календаре.")
                else:
                    for event in events:
                        await delete_calendar_event_by_id(event["id"], user_id=user_id)
                    await safe_reply_text(reply_target, f"Удалено событий: {len(events)}")
                return

            events = await find_calendar_events(parsed, user_id=user_id)
            if not events:
                await safe_reply_text(reply_target, "Не нашел подходящих событий в календаре.")
            elif len(events) == 1:
                event = events[0]
                await delete_calendar_event_by_id(event["id"], user_id=user_id)
                await safe_reply_text(reply_target, f"Событие удалено: {calendar_event_label(event)}")
            else:
                await ask_user_to_select_event(
                    reply_target,
                    chat_id,
                    "delete",
                    events,
                    parsed,
                    "Нашел несколько событий. Какое удалить?",
                )
        elif parsed["kind"] == "reschedule_calendar_event":
            delete_shape = {"delete_event": parsed["reschedule_event"]}
            events = await find_calendar_events(delete_shape, user_id=user_id)
            if not events:
                await safe_reply_text(reply_target, "Не нашел подходящих событий в календаре.")
            elif len(events) == 1:
                updated = await reschedule_calendar_event_by_payload(events[0], parsed, user_id=user_id)
                await safe_reply_text(reply_target, f"Событие перенесено: {calendar_event_label(updated)}")
            else:
                await ask_user_to_select_event(
                    reply_target,
                    chat_id,
                    "reschedule",
                    events,
                    parsed,
                    "Нашел несколько событий. Какое перенести?",
                )
        elif parsed["kind"] == "delete_reminders":
            await execute_delete_reminders(parsed, chat_id, reply_target)
        else:
            await safe_reply_text(reply_target, "Неизвестный тип действия.")
    except Exception as exc:
        logger.exception("Action execution failed")
        await safe_reply_text(reply_target, f"Не получилось выполнить действие: {exc}")

async def ask_user_to_select_event(
    reply_target: Any,
    chat_id: int,
    action: str,
    events: list[dict[str, Any]],
    payload: dict[str, Any],
    title: str,
) -> None:
    buttons = []
    lines = [title]
    for index, event in enumerate(events[:5]):
        selection_id = store.add_event_selection(chat_id, action, event, payload)
        lines.append(f"{index + 1}. {calendar_event_label(event)}")
        buttons.append(
            [InlineKeyboardButton(f"{index + 1}", callback_data=f"event_select:{selection_id}")]
        )

    await safe_reply_text(
        reply_target,
        "\n".join(lines),
        reply_markup=InlineKeyboardMarkup(buttons),
    )


def reminder_label(row: sqlite3.Row) -> str:
    return f"{format_msk_dt(row['remind_at'])} — {row['text']}"


def remove_scheduled_reminder_jobs(reminder_ids: list[int]) -> None:
    for reminder_id in reminder_ids:
        try:
            scheduler.remove_job(f"reminder:{reminder_id}")
        except Exception:
            logger.debug("Reminder job %s was not scheduled or already removed", reminder_id)


async def show_delete_all_reminders_menu(reply_target: Any) -> None:
    keyboard = InlineKeyboardMarkup(
        [
            [
                InlineKeyboardButton("Сегодня", callback_data="reminder_delete_scope:today"),
                InlineKeyboardButton("Завтра", callback_data="reminder_delete_scope:tomorrow"),
            ],
            [
                InlineKeyboardButton("Ввести дату", callback_data="reminder_delete_scope:date"),
                InlineKeyboardButton("Все", callback_data="reminder_delete_scope:all"),
            ],
        ]
    )
    await safe_reply_text(
        reply_target,
        "Какие Telegram-напоминания удалить?",
        reply_markup=keyboard,
    )


async def delete_reminders_by_scope(
    chat_id: int,
    reply_target: Any,
    *,
    start: datetime | None = None,
    end: datetime | None = None,
    label: str = "выбранный период",
) -> None:
    rows = store.search_reminders(chat_id, None, start, end, limit=10000)
    if not rows:
        await safe_reply_text(reply_target, f"Не нашел Telegram-напоминаний: {label}.")
        return

    reminder_ids = [int(row["id"]) for row in rows]
    deleted = store.cancel_reminders(chat_id, reminder_ids)
    remove_scheduled_reminder_jobs(reminder_ids)
    export_reminders_markdown(chat_id)
    await safe_reply_text(reply_target, f"Удалено Telegram-напоминаний ({label}): {deleted}")


async def execute_delete_reminders(
    action: dict[str, Any],
    chat_id: int,
    reply_target: Any,
) -> None:
    payload = action["delete_reminders"]
    start = parse_dt(payload["time_min"]) if payload.get("time_min") else None
    end = parse_dt(payload["time_max"]) if payload.get("time_max") else None
    query = payload.get("query") or None
    if payload.get("delete_all") and not query and start is None and end is None:
        await show_delete_all_reminders_menu(reply_target)
        return

    rows = store.search_reminders(
        chat_id,
        query,
        start,
        end,
        limit=10000 if payload.get("delete_all") else 100,
    )

    if not rows:
        await safe_reply_text(reply_target, "Не нашел подходящих Telegram-напоминаний.")
        return

    reminder_ids = [int(row["id"]) for row in rows]
    if payload.get("delete_all") or len(rows) == 1:
        ids_to_delete = reminder_ids if payload.get("delete_all") else [reminder_ids[0]]
        deleted = store.cancel_reminders(chat_id, ids_to_delete)
        remove_scheduled_reminder_jobs(ids_to_delete)
        export_reminders_markdown(chat_id)
        await safe_reply_text(reply_target, f"Удалено Telegram-напоминаний: {deleted}")
        return

    lines = ["Нашел несколько напоминаний. Что удалить?"]
    buttons: list[list[InlineKeyboardButton]] = []
    for index, row in enumerate(rows[:10], start=1):
        selection_id = store.add_reminder_selection(chat_id, [int(row["id"])])
        lines.append(f"{index}. {reminder_label(row)}")
        buttons.append([InlineKeyboardButton(str(index), callback_data=f"reminder_select:{selection_id}")])
    all_selection_id = store.add_reminder_selection(chat_id, reminder_ids)
    buttons.append([InlineKeyboardButton("Удалить все найденные", callback_data=f"reminder_select:{all_selection_id}")])

    await safe_reply_text(
        reply_target,
        "\n".join(lines),
        reply_markup=InlineKeyboardMarkup(buttons),
    )


async def reschedule_calendar_event_by_payload(
    event: dict[str, Any],
    payload: dict[str, Any],
    user_id: int | None = None,
) -> dict[str, Any]:
    reschedule = payload["reschedule_event"]
    new_start = parse_dt(reschedule["new_start"])
    if reschedule.get("new_end"):
        new_end = parse_dt(reschedule["new_end"])
    else:
        old_start_raw = event.get("start", {}).get("dateTime") or event.get("start", {}).get("date")
        old_end_raw = event.get("end", {}).get("dateTime") or event.get("end", {}).get("date")
        if old_start_raw and old_end_raw:
            duration = parse_dt(old_end_raw) - parse_dt(old_start_raw)
        else:
            duration = timedelta(hours=1)
        new_end = new_start + duration

    return await update_calendar_event_time(event["id"], new_start, new_end, user_id=user_id)


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not update.message:
        return
    if not is_allowed(update):
        user_id = update.effective_user.id if update.effective_user else "unknown"
        logger.warning("Blocked /start from user_id=%s", user_id)
        await update.message.reply_text("Доступ к этому боту ограничен.")
        return

    await update.message.reply_text(
        "Готов. Можно писать текстом или голосом: создам событие, поставлю напоминание или помогу разобрать задачу."
    )


async def help_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not update.message:
        return
    if not is_allowed(update):
        await update.message.reply_text("Доступ к этому боту ограничен.")
        return

    await update.message.reply_text(
        capabilities_text(resolve_help_section(context.args)),
        reply_markup=capabilities_keyboard(),
    )


async def whoami(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not update.message or not update.effective_user:
        return

    await update.message.reply_text(
        "\n".join(
            [
                f"Ваш Telegram user id: {update.effective_user.id}",
                f"Chat id: {update.effective_chat.id if update.effective_chat else 'unknown'}",
            ]
        )
    )


async def notes_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not update.message or not update.effective_chat:
        return
    if not is_allowed(update):
        await update.message.reply_text("Доступ к этому боту ограничен.")
        return

    reply = await build_notes_reply(
        chat_id=update.effective_chat.id,
        user_id=update.effective_user.id if update.effective_user else None,
        args=context.args,
        search_notes_contextual=search_notes_contextual,
        iter_notes=iter_notes,
        format_note_results=format_note_results,
        format_msk_dt=format_msk_dt,
    )
    await safe_reply_text(update.message, reply)


async def knowledge_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not update.message or not update.effective_chat:
        return
    if not is_allowed(update):
        await update.message.reply_text("Доступ к этому боту ограничен.")
        return

    reply = await build_knowledge_reply(
        chat_id=update.effective_chat.id,
        user_id=update.effective_user.id if update.effective_user else None,
        args=context.args,
        empty_message="База знаний пока пустая. Пришли .txt/.md/.csv/.json/.pdf/.docx файлом в Telegram.",
        search_knowledge_contextual=search_knowledge_contextual,
        iter_knowledge_chunks=iter_knowledge_chunks,
        format_knowledge_results=format_knowledge_results,
        format_msk_dt=format_msk_dt,
    )
    await safe_reply_text(update.message, reply)


async def memory_status(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not update.message or not update.effective_chat:
        return
    if not is_allowed(update):
        await update.message.reply_text("Доступ к этому боту ограничен.")
        return

    count = store.count_memory_messages(update.effective_chat.id)
    await update.message.reply_text(
        "\n".join(
            [
                f"Память: {'включена' if MEMORY_ENABLED else 'выключена'}",
                f"Сообщений в памяти этого чата: {count}",
                f"Хранение: {MEMORY_RETENTION_DAYS} дн.",
                f"В prompt попадает последних сообщений: {MEMORY_RECENT_MESSAGES}",
            ]
        )
    )


async def forget_today(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not update.message or not update.effective_chat:
        return
    if not is_allowed(update):
        await update.message.reply_text("Доступ к этому боту ограничен.")
        return

    deleted = store.forget_today(update.effective_chat.id)
    await update.message.reply_text(f"Удалено сообщений из памяти за сегодня: {deleted}")


async def show_calendar_range(
    update: Update,
    title: str,
    start: datetime,
    end: datetime,
) -> None:
    if not update.message:
        return
    if not is_allowed(update):
        await update.message.reply_text("Доступ к этому боту ограничен.")
        return

    try:
        user_id = update.effective_user.id if update.effective_user else None
        events = await list_calendar_events(start, end, user_id=user_id)
    except Exception as exc:
        logger.exception("Calendar listing failed")
        await update.message.reply_text(f"Не получилось прочитать календарь: {exc}")
        return

    await safe_reply_text(update.message, format_calendar_events(events, title))


def format_reminders_for_range(chat_id: int, title: str, start: datetime, end: datetime) -> str:
    rows = store.list_reminders_between(chat_id, start, end)
    if not rows:
        return f"{title}\nНапоминаний нет."

    lines = [title]
    for index, row in enumerate(rows, start=1):
        lines.append(
            f"{index}. {format_msk_dt(row['remind_at'], include_date=False)} — {row['text']}"
        )
    return "\n".join(lines)


def render_reminders_markdown(chat_id: int, days: int = 14) -> str:
    zone = ZoneInfo(TIMEZONE)
    today = datetime.now(zone).replace(hour=0, minute=0, second=0, microsecond=0)
    rows = store.list_pending_reminders_for_chat(chat_id)
    by_date: dict[str, list[sqlite3.Row]] = {}
    for row in rows:
        remind_at = to_local_dt(row["remind_at"])
        if remind_at < today or remind_at >= today + timedelta(days=days):
            continue
        key = remind_at.strftime("%Y-%m-%d")
        by_date.setdefault(key, []).append(row)

    lines = [
        "# Список дел и напоминаний",
        "",
        f"Обновлено: {format_msk_dt(now_iso())}",
        f"Период: {today.strftime('%d.%m.%Y')} - {(today + timedelta(days=days - 1)).strftime('%d.%m.%Y')}",
        "",
    ]
    for offset in range(days):
        day = today + timedelta(days=offset)
        key = day.strftime("%Y-%m-%d")
        title = day.strftime("%d.%m.%Y")
        if offset == 0:
            title += " сегодня"
        elif offset == 1:
            title += " завтра"
        elif offset == 2:
            title += " послезавтра"
        lines.append(f"## {title}")
        day_rows = by_date.get(key, [])
        if not day_rows:
            lines.append("- Нет напоминаний")
        else:
            for row in day_rows:
                lines.append(f"- {format_msk_dt(row['remind_at'], include_date=False)} — {row['text']}")
        lines.append("")

    return "\n".join(lines).rstrip() + "\n"


def export_reminders_markdown(chat_id: int) -> None:
    REMINDERS_MD_PATH.parent.mkdir(parents=True, exist_ok=True)
    REMINDERS_MD_PATH.write_text(render_reminders_markdown(chat_id), encoding="utf-8")


async def reminders(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not update.message or not update.effective_chat:
        return
    if not is_allowed(update):
        await update.message.reply_text("Доступ к этому боту ограничен.")
        return

    reply = build_reminders_reply(
        chat_id=update.effective_chat.id,
        args=context.args,
        now=datetime.now(ZoneInfo(TIMEZONE)),
        reminder_range_from_text=reminder_range_from_text,
        day_bounds=day_bounds,
        export_reminders_markdown=export_reminders_markdown,
        format_reminders_for_range=format_reminders_for_range,
    )
    await safe_reply_text(update.message, reply)


async def today(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    zone = ZoneInfo(TIMEZONE)
    start = datetime.now(zone).replace(hour=0, minute=0, second=0, microsecond=0)
    end = start + timedelta(days=1)
    await show_calendar_range(update, "События на сегодня:", start, end)


async def tomorrow(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    zone = ZoneInfo(TIMEZONE)
    start = (datetime.now(zone) + timedelta(days=1)).replace(
        hour=0,
        minute=0,
        second=0,
        microsecond=0,
    )
    end = start + timedelta(days=1)
    await show_calendar_range(update, "События на завтра:", start, end)


async def week(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    zone = ZoneInfo(TIMEZONE)
    start = datetime.now(zone).replace(hour=0, minute=0, second=0, microsecond=0)
    end = start + timedelta(days=7)
    await show_calendar_range(update, "События на ближайшие 7 дней:", start, end)


async def actions(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not update.message or not update.effective_chat:
        return
    if not is_allowed(update):
        await update.message.reply_text("Доступ к этому боту ограничен.")
        return

    rows = store.recent_actions_for_display(update.effective_chat.id, 10)
    if not rows:
        await update.message.reply_text("Локальный журнал действий пока пуст.")
        return

    lines = ["Последние действия:"]
    for index, row in enumerate(rows, start=1):
        try:
            summary = decrypt_text(row["summary_cipher"])
        except Exception:
            summary = "[не удалось расшифровать]"
        first_line = summary.splitlines()[0] if summary else row["kind"]
        lines.append(f"{index}. {row['created_at']} — {row['kind']} — {first_line}")

    await safe_reply_text(update.message, "\n".join(lines))


async def last_action(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not update.message or not update.effective_chat:
        return
    if not is_allowed(update):
        await update.message.reply_text("Доступ к этому боту ограничен.")
        return

    rows = store.recent_actions_for_display(update.effective_chat.id, 1)
    if not rows:
        await update.message.reply_text("Последних действий нет.")
        return

    row = rows[0]
    try:
        summary = decrypt_text(row["summary_cipher"])
    except Exception:
        summary = "[не удалось расшифровать]"
    await safe_reply_text(update.message, f"Последнее действие:\n{row['created_at']} — {row['kind']}\n{summary}")


async def calendar_auth_status(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not update.message:
        return
    if not is_allowed(update):
        await update.message.reply_text("Доступ к этому боту ограничен.")
        return

    try:
        user_id = update.effective_user.id if update.effective_user else None
        await asyncio.to_thread(get_calendar_service, user_id)
        calendar_id = store.get_user_calendar_id(user_id) if user_id is not None else GOOGLE_CALENDAR_ID
    except Exception as exc:
        await safe_reply_text(
            update.message,
            "\n".join(
                [
                    "Google Calendar сейчас не авторизован для вашего Telegram user id.",
                    f"Причина: {exc}",
                    "Запусти в терминале:",
                    f"python google_calendar_setup.py --user-id {update.effective_user.id if update.effective_user else '<id>'}",
                    "Потом перезапусти бота.",
                ]
            ),
        )
        return

    await safe_reply_text(update.message, f"Google Calendar авторизация работает.\nCalendar ID: {calendar_id}")


async def calendar_set(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not update.message or not update.effective_user:
        return
    if not is_allowed(update):
        await update.message.reply_text("Доступ к этому боту ограничен.")
        return

    if not context.args:
        await safe_reply_text(
            update.message,
            "\n".join(
                [
                    f"Текущий Calendar ID: {store.get_user_calendar_id(update.effective_user.id)}",
                    "Чтобы изменить:",
                    "/calendar_set primary",
                    "или",
                    "/calendar_set your_calendar_id@group.calendar.google.com",
                ]
            ),
        )
        return

    calendar_id = " ".join(context.args).strip()
    store.set_user_calendar_id(update.effective_user.id, calendar_id)
    await safe_reply_text(update.message, f"Calendar ID сохранен для вас: {calendar_id}")


def short_debug_text(value: str, limit: int = 180) -> str:
    value = " ".join(value.split())
    if len(value) <= limit:
        return value
    return value[: limit - 1].rstrip() + "…"


async def debug_log(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not update.message or not update.effective_chat:
        return
    if not is_allowed(update):
        await update.message.reply_text("Доступ к этому боту ограничен.")
        return

    limit = 5
    if context.args and context.args[0].isdigit():
        limit = max(1, min(20, int(context.args[0])))

    rows = store.recent_debug_logs(update.effective_chat.id, limit)
    if not rows:
        await update.message.reply_text("Диагностический журнал пока пуст.")
        return

    lines = ["Последняя диагностика:"]
    for row in rows:
        try:
            text = decrypt_text(row["input_cipher"])
        except Exception:
            text = "[не удалось расшифровать]"
        lines.extend(
            [
                "",
                f"#{row['id']} {row['created_at']}",
                f"Источник: {row['source']}"
                + (f" / STT: {row['stt_provider']}" if row["stt_provider"] else ""),
                f"Текст: {short_debug_text(text)}",
                f"Intent: {row['intent_kind'] or '-'}",
                f"Результат: {row['result'] or '-'}",
            ]
        )
        if row["error"]:
            lines.append(f"Ошибка: {short_debug_text(row['error'], 120)}")

    await safe_reply_text(update.message, "\n".join(lines))


async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not is_allowed(update):
        user_id = update.effective_user.id if update.effective_user else "unknown"
        logger.warning("Blocked text message from user_id=%s", user_id)
        if update.message:
            await update.message.reply_text("Доступ к этому боту ограничен.")
        return

    if update.message and update.message.text:
        await handle_text(update.message.text, update, context)


async def handle_voice(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not is_allowed(update):
        user_id = update.effective_user.id if update.effective_user else "unknown"
        logger.warning("Blocked voice message from user_id=%s", user_id)
        if update.message:
            await update.message.reply_text("Доступ к этому боту ограничен.")
        return

    if not update.message:
        return

    debug_id: int | None = None
    try:
        text, stt_provider = await transcribe_voice(update, context)
        debug_id = store.add_debug_log(
            chat_id=update.effective_chat.id,
            user_id=update.effective_user.id if update.effective_user else None,
            source="voice",
            stt_provider=stt_provider,
            input_text=text,
        )
    except ImportError:
        await update.message.reply_text("Для голосовых нужно установить `faster-whisper`.")
        return
    except ValueError as exc:
        await update.message.reply_text(str(exc))
        return
    except Exception:
        logger.exception("Voice transcription failed")
        await update.message.reply_text("Не смог распознать голосовое сообщение.")
        return

    if not text:
        await update.message.reply_text("Не услышал текст в голосовом.")
        return

    await safe_reply_text(update.message, f"Распознал ({stt_provider}): {text}")
    await handle_text(
        text,
        update,
        context,
        source="voice",
        stt_provider=stt_provider,
        debug_id=debug_id,
    )


async def handle_document(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not is_allowed(update):
        user_id = update.effective_user.id if update.effective_user else "unknown"
        logger.warning("Blocked document message from user_id=%s", user_id)
        if update.message:
            await update.message.reply_text("Доступ к этому боту ограничен.")
        return

    if not update.message or not update.message.document or not update.effective_chat:
        return

    document = update.message.document
    filename = document.file_name or "document"
    suffix = Path(filename).suffix.lower()
    if suffix not in supported_knowledge_suffixes():
        await safe_reply_text(
            update.message,
            "Пока умею добавлять в базу знаний: txt, md, csv, json, yaml, log, html, xml, pdf, docx.",
        )
        return

    max_bytes = KNOWLEDGE_MAX_FILE_MB * 1024 * 1024
    if document.file_size and document.file_size > max_bytes:
        await safe_reply_text(
            update.message,
            f"Файл слишком большой: лимит {KNOWLEDGE_MAX_FILE_MB} МБ.",
        )
        return

    await safe_send_typing(context, update.effective_chat.id)
    try:
        telegram_file = await context.bot.get_file(document.file_id)
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / filename
            await telegram_file.download_to_drive(custom_path=str(path))
            text = await asyncio.to_thread(extract_text_from_file, path, filename)
        info = add_knowledge_document(
            update.effective_chat.id,
            update.effective_user.id if update.effective_user else None,
            filename,
            text,
        )
    except Exception as exc:
        logger.exception("Knowledge file indexing failed")
        await safe_reply_text(update.message, f"Не получилось добавить файл в базу знаний: {exc}")
        return

    reply = "\n".join(
        [
            "Файл добавлен в базу знаний.",
            f"Название: {info['filename']}",
            f"Фрагментов: {info['chunk_count']}",
            "Теперь можно спросить: /kb <что найти>",
        ]
    )
    store.add_action_history(
        update.effective_chat.id,
        "knowledge_add",
        reply,
        {"document_id": info["document_id"], "filename": filename},
    )
    await safe_reply_text(update.message, reply)


async def handle_confirmation(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    query = update.callback_query
    if not query:
        return

    await query.answer()
    if not is_allowed(update):
        user_id = update.effective_user.id if update.effective_user else "unknown"
        logger.warning("Blocked callback from user_id=%s", user_id)
        await query.edit_message_text("Доступ к этому боту ограничен.")
        return

    action, raw_id = query.data.split(":", 1)
    action_id = int(raw_id)

    row = store.pop_pending(action_id)
    if not row:
        await query.edit_message_text("Это действие уже обработано или устарело.")
        return

    if action == "cancel":
        await query.edit_message_text("Отменено.")
        return

    parsed = json.loads(row["action_json"])
    user_id = int(row["user_id"]) if row["user_id"] is not None else None

    try:
        if parsed["kind"] == "calendar_event":
            link = await create_calendar_event(parsed, user_id=user_id)
            await query.edit_message_text(calendar_event_summary(parsed, link))
        elif parsed["kind"] == "reminder":
            reminder = parsed["reminder"]
            reminder_id = store.add_reminder(
                chat_id=row["chat_id"],
                text=reminder["text"],
                remind_at=reminder["remind_at"],
            )
            await schedule_reminder(
                context.application,
                reminder_id,
                row["chat_id"],
                reminder["text"],
                reminder["remind_at"],
            )
            export_reminders_markdown(row["chat_id"])
            await query.edit_message_text(
                "\n".join(
                    [
                        "Напоминание поставлено.",
                        f"Когда: {format_msk_dt(reminder['remind_at'])}",
                        f"Текст: {reminder['text']}",
                    ]
                )
            )
        elif parsed["kind"] == "recurring_reminders":
            recurring = parsed["recurring_reminders"]
            reminder_texts = recurring.get("texts") or []
            for index, remind_at in enumerate(recurring["occurrences"]):
                reminder_text = reminder_texts[index] if index < len(reminder_texts) else recurring["text"]
                reminder_id = store.add_reminder(
                    chat_id=row["chat_id"],
                    text=reminder_text,
                    remind_at=remind_at,
                )
                await schedule_reminder(
                    context.application,
                    reminder_id,
                    row["chat_id"],
                    reminder_text,
                    remind_at,
                )
            export_reminders_markdown(row["chat_id"])
            await query.edit_message_text(
                "\n".join(
                    [
                        "Повторяющиеся напоминания поставлены.",
                        f"Количество: {recurring['count']}",
                        f"Первое: {format_msk_dt(recurring['occurrences'][0])}",
                        f"Последнее: {format_msk_dt(recurring['occurrences'][-1])}",
                    ]
                )
            )
        elif parsed["kind"] == "recurring_calendar_events":
            recurring = parsed["recurring_calendar_events"]
            for occurrence in recurring["occurrences"]:
                await create_calendar_event_from_fields(
                    title=recurring["title"],
                    start=occurrence["start"],
                    end=occurrence["end"],
                    description=recurring.get("description", ""),
                    reminder_minutes=recurring.get("reminder_minutes"),
                    user_id=user_id,
                )
            await query.edit_message_text(
                "\n".join(
                    [
                        "Повторяющиеся события созданы.",
                        f"Количество: {recurring['count']}",
                        f"Первое: {format_msk_dt(recurring['occurrences'][0]['start'])}",
                        f"Последнее: {format_msk_dt(recurring['occurrences'][-1]['start'])}",
                        f"Напоминание: за {recurring['reminder_minutes']} мин.",
                    ]
                )
            )
        elif parsed["kind"] == "delete_calendar_event":
            if parsed["delete_event"].get("delete_all"):
                events = await find_calendar_events_limited(parsed, limit=100, user_id=user_id)
                if not events:
                    await query.edit_message_text("Не нашел подходящих событий в календаре.")
                else:
                    for event in events:
                        await delete_calendar_event_by_id(event["id"], user_id=user_id)
                    await query.edit_message_text(f"Удалено событий: {len(events)}")
                return

            events = await find_calendar_events(parsed, user_id=user_id)
            if not events:
                await query.edit_message_text("Не нашел подходящих событий в календаре.")
            elif len(events) == 1:
                event = events[0]
                await delete_calendar_event_by_id(event["id"], user_id=user_id)
                await query.edit_message_text(f"Событие удалено: {calendar_event_label(event)}")
            else:
                await query.edit_message_text("Нашел несколько событий. Выберите ниже.")
                await ask_user_to_select_event(
                    query.message,
                    int(row["chat_id"]),
                    "delete",
                    events,
                    parsed,
                    "Какое событие удалить?",
                )
        elif parsed["kind"] == "reschedule_calendar_event":
            delete_shape = {"delete_event": parsed["reschedule_event"]}
            events = await find_calendar_events(delete_shape, user_id=user_id)
            if not events:
                await query.edit_message_text("Не нашел подходящих событий в календаре.")
            elif len(events) == 1:
                updated = await reschedule_calendar_event_by_payload(events[0], parsed, user_id=user_id)
                await query.edit_message_text(f"Событие перенесено: {calendar_event_label(updated)}")
            else:
                await query.edit_message_text("Нашел несколько событий. Выберите ниже.")
                await ask_user_to_select_event(
                    query.message,
                    int(row["chat_id"]),
                    "reschedule",
                    events,
                    parsed,
                    "Какое событие перенести?",
                )
        elif parsed["kind"] == "delete_reminders":
            await execute_delete_reminders(parsed, int(row["chat_id"]), query.message)
        else:
            await query.edit_message_text("Неизвестный тип действия.")
    except Exception as exc:
        logger.exception("Action execution failed")
        await query.edit_message_text(f"Не получилось выполнить действие: {exc}")


async def handle_event_selection(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    query = update.callback_query
    if not query:
        return

    await query.answer()
    if not is_allowed(update):
        await query.edit_message_text("Доступ к этому боту ограничен.")
        return

    _prefix, raw_id = query.data.split(":", 1)
    row = store.pop_event_selection(int(raw_id))
    if not row:
        await query.edit_message_text("Этот выбор уже обработан или устарел.")
        return

    event = json.loads(row["event_json"])
    payload = json.loads(row["payload_json"])
    action = row["action"]

    try:
        user_id = update.effective_user.id if update.effective_user else None
        if action == "delete":
            await delete_calendar_event_by_id(event["id"], user_id=user_id)
            await query.edit_message_text(f"Событие удалено: {calendar_event_label(event)}")
        elif action == "reschedule":
            updated = await reschedule_calendar_event_by_payload(event, payload, user_id=user_id)
            await query.edit_message_text(f"Событие перенесено: {calendar_event_label(updated)}")
        else:
            await query.edit_message_text("Неизвестный выбор.")
    except Exception as exc:
        logger.exception("Selected event action failed")
        await query.edit_message_text(f"Не получилось выполнить действие: {exc}")


async def handle_reminder_selection(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    query = update.callback_query
    if not query:
        return

    await query.answer()
    if not is_allowed(update):
        await query.edit_message_text("Доступ к этому боту ограничен.")
        return

    _prefix, raw_id = query.data.split(":", 1)
    row = store.pop_reminder_selection(int(raw_id))
    if not row:
        await query.edit_message_text("Этот выбор уже обработан или устарел.")
        return

    reminder_ids = json.loads(row["reminder_ids_json"])
    try:
        ids_to_delete = [int(value) for value in reminder_ids]
        deleted = store.cancel_reminders(int(row["chat_id"]), ids_to_delete)
        remove_scheduled_reminder_jobs(ids_to_delete)
        export_reminders_markdown(int(row["chat_id"]))
        await query.edit_message_text(f"Удалено Telegram-напоминаний: {deleted}")
    except Exception as exc:
        logger.exception("Selected reminder delete failed")
        await query.edit_message_text(f"Не получилось удалить напоминания: {exc}")


async def handle_reminder_delete_scope(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    query = update.callback_query
    if not query:
        return

    await query.answer()
    if not is_allowed(update):
        await query.edit_message_text("Доступ к этому боту ограничен.")
        return

    _prefix, scope = query.data.split(":", 1)
    chat_id = update.effective_chat.id if update.effective_chat else query.message.chat_id
    zone = ZoneInfo(TIMEZONE)
    today = datetime.now(zone).replace(hour=0, minute=0, second=0, microsecond=0)

    if scope == "date":
        await query.edit_message_text(
            "Напиши дату текстом, например:\n"
            "удали все напоминания на 25 мая\n"
            "или\n"
            "удали все напоминания на 25.05"
        )
        return

    await query.edit_message_text("Удаляю...")
    if scope == "today":
        start, end = today, today + timedelta(days=1)
        await delete_reminders_by_scope(chat_id, query.message, start=start, end=end, label="сегодня")
    elif scope == "tomorrow":
        start, end = today + timedelta(days=1), today + timedelta(days=2)
        await delete_reminders_by_scope(chat_id, query.message, start=start, end=end, label="завтра")
    elif scope == "all":
        await delete_reminders_by_scope(chat_id, query.message, label="все")
    else:
        await safe_reply_text(query.message, "Неизвестный вариант удаления.")


async def handle_help_callback(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    query = update.callback_query
    if not query:
        return

    await query.answer()
    if not is_allowed(update):
        await query.edit_message_text("Доступ к этому боту ограничен.")
        return

    _prefix, section = query.data.split(":", 1)
    await query.edit_message_text(
        capabilities_text(section),
        reply_markup=capabilities_keyboard(),
    )


async def privacy_status(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not update.message:
        return
    if not is_allowed(update):
        await update.message.reply_text("Доступ к этому боту ограничен.")
        return

    lines = [
        f"Режим приватности: {PRIVACY_MODE}",
        f"STT-провайдеры: {', '.join(selected_stt_providers())}",
        f"Лимит голосового: {STT_MAX_VOICE_SECONDS} сек.",
    ]
    if PRIVACY_MODE == "strict":
        lines.append("Во внешние STT-сервисы голос не отправляется.")
    else:
        lines.append("Голос может уходить выбранному STT-провайдеру; при ошибке включается fallback.")
    await safe_reply_text(update.message, "\n".join(lines))


async def stt_status(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not update.message:
        return
    if not is_allowed(update):
        await update.message.reply_text("Доступ к этому боту ограничен.")
        return

    lines = [
        f"Основной STT: {STT_PROVIDER}",
        f"Fallback STT: {STT_FALLBACK_PROVIDER}",
        f"Активная цепочка: {', '.join(selected_stt_providers())}",
        f"Локальный Whisper: {WHISPER_MODEL}",
        f"Groq model: {GROQ_STT_MODEL}",
        f"OpenAI STT model: {OPENAI_STT_MODEL}",
        f"Groq key: {'есть' if GROQ_API_KEY else 'нет'}",
        f"OpenAI key: {'есть' if OPENAI_API_KEY else 'нет'}",
    ]
    await safe_reply_text(update.message, "\n".join(lines))


async def post_init(app: Application) -> None:
    for row in store.list_pending_reminders():
        remind_at = parse_dt(row["remind_at"])
        if remind_at <= datetime.now(ZoneInfo(TIMEZONE)):
            store.mark_reminder_sent(row["id"])
            continue
        await schedule_reminder(
            app,
            int(row["id"]),
            int(row["chat_id"]),
            str(row["text"]),
            str(row["remind_at"]),
        )
    scheduler.start()


def main() -> None:
    if not TELEGRAM_BOT_TOKEN:
        raise RuntimeError("Set TELEGRAM_BOT_TOKEN in .env")
    if "PASTE_YOUR_BOTFATHER_TOKEN_HERE" in TELEGRAM_BOT_TOKEN:
        raise RuntimeError("Replace TELEGRAM_BOT_TOKEN in .env with the real BotFather token")

    asyncio.set_event_loop(asyncio.new_event_loop())

    app = Application.builder().token(TELEGRAM_BOT_TOKEN).post_init(post_init).build()
    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("help", help_command))
    app.add_handler(CommandHandler("whoami", whoami))
    app.add_handler(CommandHandler("notes", notes_command))
    app.add_handler(CommandHandler("kb", knowledge_command))
    app.add_handler(CommandHandler("knowledge", knowledge_command))
    app.add_handler(CommandHandler("memory", memory_status))
    app.add_handler(CommandHandler("forget_today", forget_today))
    app.add_handler(CommandHandler("today", today))
    app.add_handler(CommandHandler("tomorrow", tomorrow))
    app.add_handler(CommandHandler("week", week))
    app.add_handler(CommandHandler("reminders", reminders))
    app.add_handler(CommandHandler("actions", actions))
    app.add_handler(CommandHandler("last", last_action))
    app.add_handler(CommandHandler("calendar_auth", calendar_auth_status))
    app.add_handler(CommandHandler("calendar_set", calendar_set))
    app.add_handler(CommandHandler("debuglog", debug_log))
    app.add_handler(CommandHandler("privacy", privacy_status))
    app.add_handler(CommandHandler("stt", stt_status))
    app.add_handler(CallbackQueryHandler(handle_help_callback, pattern=r"^help:(calendar|reminders|voice|memory|notes|knowledge)$"))
    app.add_handler(CallbackQueryHandler(handle_event_selection, pattern=r"^event_select:\d+$"))
    app.add_handler(CallbackQueryHandler(handle_reminder_selection, pattern=r"^reminder_select:\d+$"))
    app.add_handler(CallbackQueryHandler(handle_reminder_delete_scope, pattern=r"^reminder_delete_scope:(today|tomorrow|date|all)$"))
    app.add_handler(CallbackQueryHandler(handle_confirmation, pattern=r"^(confirm|cancel):\d+$"))
    app.add_handler(MessageHandler(filters.VOICE, handle_voice))
    app.add_handler(MessageHandler(filters.Document.ALL, handle_document))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_message))

    logger.info("Telegram local organizer bot started")
    app.run_polling(allowed_updates=Update.ALL_TYPES)


if __name__ == "__main__":
    main()
